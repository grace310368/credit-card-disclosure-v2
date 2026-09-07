from __future__ import annotations

import argparse
import io
import json
import re
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# 避免在唯讀執行沙箱中產生 __pycache__/*.pyc（會觸發 Refusing to overwrite）
sys.dont_write_bytecode = True

from percent_utils import PERCENT_DECIMAL_FIELDS, apply_percent_number_format, normalize_percent_value
from bank_aliases import ITEM_RANKS
from workbook_block_helpers import apply_default_font

DEFAULT_WORKBOOK = Path('銀行局信用卡公開資料.xlsx')
ZIP_SUFFIX = '信用卡重要資訊揭露.zip'
DEFAULT_ZIP_SEARCH_DIRS = [Path('.'), Path('input')]
MAIN_SHEET = '歷史資料(年+月)'
MARKET_TOTAL_ITEM = '市場總計(銀行局)'
BANK_BUREAU_ITEM = '銀行局'
JCIC_ITEM = '財團法人金融聯合徵信中心'
SPECIAL_ITEMS = [MARKET_TOTAL_ITEM, BANK_BUREAU_ITEM, JCIC_ITEM]

LONGFORM_FIELD_ALIASES = {
    'yyyymm': ['YYYYMM'],
    'ad_year': ['年度'],
    'month_number': ['月份'],
    'rank': ['Rank'],
    'item': ['Item'],
    'circulating_cards': ['流通卡數'],
    'valid_cards': ['有效卡數'],
    'new_cards_this_month': ['當月發卡數'],
    'cancelled_cards_this_month': ['當月停卡數'],
    'revolving_balance_million': ['循環信用餘額'],
    'installment_balance_not_yet_due_million': ['未到期分期付款餘額'],
    'signed_amount_million': ['當月簽帳金額'],
    'cash_advance_amount_million': ['當月預借現金金額'],
    'overdue_3m_ratio_percent': [
        '逾期三個月以上帳款占應收帳款餘額（含催收款）之比率(%)',
        '逾期三個月以上帳款占應收帳款餘額(含催收款)之比率(%)',
    ],
    'overdue_6m_ratio_percent': [
        '逾期六個月以上帳款占應收帳款餘額（含催收款）之比率(%)',
        '逾期六個月以上帳款占應收帳款餘額(含催收款)之比率(%)',
    ],
    'allowance_coverage_ratio_percent': ['備抵呆帳提足率(%)', '備抵呆帳提足率'],
    'charge_off_amount_this_month_million': ['當月轉銷呆帳金額', '當月轉銷呆帳金額'],
    'charge_off_amount_ytd_million': ['當年度累計轉銷呆帳金額', '當年度轉銷呆帳金額累計至資料月份', '當年度累計轉銷呆帳金額'],
    'market_total': ['市場總計'],
}

PERCENT_DISPLAY_FIELDS = PERCENT_DECIMAL_FIELDS

LONGFORM_MARKET_FIELDS = [
    'circulating_cards',
    'valid_cards',
    'new_cards_this_month',
    'cancelled_cards_this_month',
    'revolving_balance_million',
    'installment_balance_not_yet_due_million',
    'signed_amount_million',
    'cash_advance_amount_million',
    'overdue_3m_ratio_percent',
    'overdue_6m_ratio_percent',
    'allowance_coverage_ratio_percent',
    'charge_off_amount_this_month_million',
    'charge_off_amount_ytd_million',
]

METRIC_LABELS = {
    'circulating_cards': '流通卡數',
    'valid_cards': '有效卡數',
    'new_cards_this_month': '當月發卡數',
    'cancelled_cards_this_month': '當月停卡數',
    'revolving_balance_million': '循環信用餘額',
    'installment_balance_not_yet_due_million': '未到期分期付款餘額',
    'signed_amount_million': '當月簽帳金額',
    'cash_advance_amount_million': '當月預借現金金額',
    'overdue_3m_ratio_percent': '逾期三個月以上帳款占應收帳款餘額(含催收款)之比率(%)',
    'overdue_6m_ratio_percent': '逾期六個月以上帳款占應收帳款餘額(含催收款)之比率(%)',
    'allowance_coverage_ratio_percent': '備抵呆帳提足率(%)',
    'charge_off_amount_this_month_million': '當月轉銷呆帳金額',
    'charge_off_amount_ytd_million': '當年度轉銷呆帳金額累計至資料月份',
    'market_total': '市場總計',
}


@dataclass
class DownloadMeta:
    mode: str | None
    ssl_fallback_used: bool
    warning: str | None = None


@dataclass
class WorkbookResolution:
    path: Path
    fallback_used: bool
    warnings: list[str]


@dataclass
class FetchOptions:
    insecure: bool
    prefer_local_zip: bool
    zip_search_dirs: list[Path]
    explicit_market: dict[str, Any] | None = None


def normalize_header(value: Any) -> str:
    return re.sub(r'\s+', '', str(value or '').replace('\u3000', ''))


def parse_roc_month(text: str) -> tuple[int, int]:
    match = re.search(r'(\d{3})\D*(\d{1,2})', str(text))
    if not match:
        raise ValueError(f'無法解析民國月份: {text}')
    return int(match.group(1)), int(match.group(2))


def roc_month_text(y: int, m: int) -> str:
    return f'{y}年{m:02d}月'


def roc_to_ad_yyyymm(y: int, m: int) -> int:
    return (y + 1911) * 100 + m


def ad_yyyymm_to_roc(yyyymm: int) -> tuple[int, int]:
    ad_y, m = divmod(int(yyyymm), 100)
    return ad_y - 1911, m


def build_zip_url(y: int, m: int) -> str:
    return f'https://www.fsc.gov.tw/userfiles/file/{y:03d}{m:02d}_{urllib.parse.quote(ZIP_SUFFIX)}'


def is_ssl_verification_error(exc: Exception) -> bool:
    if isinstance(exc, ssl.SSLCertVerificationError):
        return True
    if isinstance(exc, ssl.SSLError):
        return 'CERTIFICATE_VERIFY_FAILED' in str(exc).upper()
    if isinstance(exc, urllib.error.URLError):
        reason = getattr(exc, 'reason', None)
        if isinstance(reason, Exception):
            return is_ssl_verification_error(reason)
    return 'CERTIFICATE_VERIFY_FAILED' in str(exc).upper()


def make_ssl_context(insecure: bool) -> ssl.SSLContext:
    ctx = ssl.create_default_context()
    if insecure:
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
    return ctx


def download_bytes(url: str, timeout: int = 60, insecure: bool = False) -> tuple[bytes, DownloadMeta]:
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})

    if insecure:
        with urllib.request.urlopen(req, context=make_ssl_context(True), timeout=timeout) as resp:
            return resp.read(), DownloadMeta('https_insecure_forced', False)

    try:
        with urllib.request.urlopen(req, context=make_ssl_context(False), timeout=timeout) as resp:
            return resp.read(), DownloadMeta('https_verified', False)
    except Exception as exc:
        if not is_ssl_verification_error(exc):
            raise
        with urllib.request.urlopen(req, context=make_ssl_context(True), timeout=timeout) as resp:
            return resp.read(), DownloadMeta(
                'https_insecure_fallback',
                True,
                f'SSL 憑證驗證失敗，已自動改用不驗證模式: {type(exc).__name__}: {exc}',
            )


def to_number(value: Any) -> float | int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value

    text = str(value).strip().replace(',', '')
    if text in {'', '-', '－', '—'}:
        return None
    try:
        num = float(text)
        return int(num) if num.is_integer() else num
    except ValueError:
        return None


def validate_xlsx(path: Path) -> str | None:
    if not path.exists():
        return f'檔案不存在: {path}'
    if path.suffix.lower() != '.xlsx':
        return f'副檔名不是 .xlsx: {path}'
    if not zipfile.is_zipfile(path):
        return f'不是有效的 xlsx/zip 檔: {path}'

    try:
        with zipfile.ZipFile(path) as zf:
            if '[Content_Types].xml' not in zf.namelist():
                return f'缺少 [Content_Types].xml: {path}'
        wb = load_workbook(path, read_only=True, data_only=True)
        wb.close()
        return None
    except Exception as exc:
        return f'{type(exc).__name__}: {exc}'


def workbook_sort_key(path: Path) -> tuple[int, int, str]:
    name = path.name.lower()
    if name == '銀行局信用卡公開資料.xlsx':
        rank = 0
    elif name.startswith('銀行局信用卡公開資料_') and 'backup' not in name:
        rank = 1
    elif name.startswith('銀行局信用卡公開資料_backup_'):
        rank = 2
    elif name.startswith('workbook'):
        rank = 3
    else:
        rank = 9

    try:
        mtime = -int(path.stat().st_mtime)
    except Exception:
        mtime = 0
    return rank, mtime, name


def list_workbook_candidates() -> list[Path]:
    patterns = ['銀行局信用卡公開資料*.xlsx', 'workbook*.xlsx', '*.xlsx']
    seen: set[Path] = set()
    out: list[Path] = []

    for pattern in patterns:
        for path in Path('.').glob(pattern):
            if path.name.startswith('~$'):
                continue
            resolved = path.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            out.append(path)

    return sorted(out, key=workbook_sort_key)


def resolve_workbook(requested: Path, strict: bool) -> WorkbookResolution:
    error = validate_xlsx(requested)
    if error is None:
        return WorkbookResolution(requested, False, [])

    warnings = [f'指定工作簿不可用: {requested} ({error})']
    if strict:
        raise FileNotFoundError(warnings[-1])

    requested_resolved = requested.resolve()
    for candidate in list_workbook_candidates():
        if candidate.resolve() == requested_resolved:
            continue
        error = validate_xlsx(candidate)
        if error is None:
            warnings.append(f'已自動改用可用工作簿: {candidate}')
            return WorkbookResolution(candidate, True, warnings)
        warnings.append(f'略過不可用工作簿: {candidate} ({error})')

    raise FileNotFoundError('找不到可用工作簿。' + ('；' + '；'.join(warnings) if warnings else ''))


def describe_non_zip_content(data: bytes) -> str | None:
    head = data[:200].lstrip().lower()
    if head.startswith(b'<!doctype html') or head.startswith(b'<html'):
        return '內容看起來是 HTML，不是 ZIP，可能是下載到錯誤頁或另存了網頁內容'
    if head.startswith(b'{') or head.startswith(b'['):
        return '內容看起來是 JSON，不是 ZIP'
    return None


def load_zip_bytes(source: Path | bytes, source_label: str) -> bytes:
    data = source if isinstance(source, bytes) else source.read_bytes()
    if zipfile.is_zipfile(io.BytesIO(data)):
        return data

    hint = describe_non_zip_content(data)
    if hint:
        raise zipfile.BadZipFile(f'{source_label}: {hint}')
    raise zipfile.BadZipFile(f'{source_label}: 不是有效的 ZIP 檔')


def scan_market_month(ws) -> tuple[int, int]:
    candidates = [ws['J3'].value]
    for row in range(1, min(ws.max_row, 10) + 1):
        for col in range(1, min(ws.max_column, 15) + 1):
            value = ws.cell(row, col).value
            if value is not None and re.search(r'\d{3}\D*\d{1,2}', str(value)):
                candidates.append(value)

    for candidate in candidates:
        try:
            return parse_roc_month(str(candidate))
        except ValueError:
            continue
    raise ValueError('無法在銀行局檔案中確認資料月份')


def extract_market_total(ws, source_label: str) -> dict[str, Any]:
    roc_y, roc_m = scan_market_month(ws)
    headers = [normalize_header(ws.cell(4, col).value) for col in range(1, ws.max_column + 1)]

    total_row = None
    for row in range(5, ws.max_row + 1):
        if normalize_header(ws.cell(row, 1).value) == '總計':
            total_row = row
            break
    if total_row is None:
        raise ValueError(f'找不到銀行局總計列: {source_label}')

    value_cells = {
        header: ws.cell(total_row, col)
        for col, header in enumerate(headers, start=1)
        if header
    }

    def metric(header_name: str) -> float | int:
        key = normalize_header(header_name)
        if key not in value_cells:
            raise KeyError(f'銀行局檔案找不到欄位: {header_name}; source={source_label}')
        cell = value_cells[key]
        number = to_number(cell.value)
        if number is None:
            raise ValueError(f'銀行局欄位無法轉數字: {header_name}={cell.value!r}; source={source_label}')
        return number

    def optional_metric(*header_names: str) -> float | int | None:
        for header_name in header_names:
            key = normalize_header(header_name)
            if key not in value_cells:
                continue
            return to_number(value_cells[key].value)
        return None

    def optional_percent_metric(*header_names: str, force_percent_input: bool = False) -> float | None:
        for header_name in header_names:
            key = normalize_header(header_name)
            if key not in value_cells:
                continue
            cell = value_cells[key]
            value = normalize_percent_value(
                cell.value,
                number_format=str(cell.number_format or ''),
                force_percent_input=force_percent_input,
            )
            if value is not None:
                return value
        return None

    def optional_metric_div_1000(*header_names: str) -> float | int | None:
        value = optional_metric(*header_names)
        return None if value is None else value / 1000

    return {
        'month': roc_month_text(roc_y, roc_m),
        'roc_year': roc_y,
        'roc_month': roc_m,
        'ad_yyyymm': roc_to_ad_yyyymm(roc_y, roc_m),
        'ad_year': roc_y + 1911,
        'month_number': roc_m,
        'circulating_cards': metric('流通卡數'),
        'valid_cards': metric('有效卡數'),
        'new_cards_this_month': metric('當月發卡數'),
        'cancelled_cards_this_month': metric('當月停卡數'),
        'signed_amount_million': metric('當月簽帳金額') / 1000,
        'revolving_balance_million': metric('循環信用餘額') / 1000,
        'installment_balance_not_yet_due_million': optional_metric_div_1000('未到期分期付款餘額'),
        'cash_advance_amount_million': optional_metric_div_1000('當月預借現金金額'),
        # 銀行局檔案的逾期比率是裸的百分比數字（0.12 代表 0.12%），一律 /100 成小數比率
        'overdue_3m_ratio_percent': optional_percent_metric('逾期三個月以上帳款占應收帳款餘額（含催收款）之比率(%)', '逾期三個月以上帳款占應收帳款餘額(含催收款)之比率(%)', force_percent_input=True),
        'overdue_6m_ratio_percent': optional_percent_metric('逾期六個月以上帳款占應收帳款餘額（含催收款）之比率(%)', '逾期六個月以上帳款占應收帳款餘額(含催收款)之比率(%)', force_percent_input=True),
        'allowance_coverage_ratio_percent': optional_percent_metric('備抵呆帳提足率(%)', '備抵呆帳提足率'),
        'charge_off_amount_this_month_million': optional_metric_div_1000('當月轉銷呆帳金額', '當月轉銷呆帳金額'),
        'charge_off_amount_ytd_million': optional_metric_div_1000('當年度累計轉銷呆帳金額', '當年度轉銷呆帳金額累計至資料月份', '當年度累計轉銷呆帳金額'),
    }


def parse_market_from_xlsx_bytes(xlsx_bytes: bytes, source_label: str) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb[wb.sheetnames[0]]
    return extract_market_total(ws, source_label)


def parse_market_from_zip_bytes(zip_bytes: bytes, source_label: str) -> dict[str, Any]:
    zip_bytes = load_zip_bytes(zip_bytes, source_label)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        xlsx_names = [name for name in zf.namelist() if name.lower().endswith('.xlsx')]
        if not xlsx_names:
            raise ValueError(f'ZIP 內找不到 xlsx: {source_label}')
        return parse_market_from_xlsx_bytes(zf.read(xlsx_names[0]), source_label)


def with_source_meta(
    market: dict[str, Any],
    *,
    source_kind: str,
    source_url: str | None,
    local_zip_path: str | None,
    download_mode: str | None,
    ssl_fallback_used: bool,
    warnings: list[str],
) -> dict[str, Any]:
    return {
        **market,
        'source_kind': source_kind,
        'source_url': source_url,
        'local_zip_path': local_zip_path,
        'download_mode': download_mode,
        'ssl_fallback_used': ssl_fallback_used,
        'warnings': warnings,
    }


def parse_market_from_zip_file(path: Path) -> dict[str, Any]:
    market = parse_market_from_zip_bytes(path.read_bytes(), str(path))
    return with_source_meta(
        market,
        source_kind='local_zip_explicit',
        source_url=None,
        local_zip_path=str(path),
        download_mode=None,
        ssl_fallback_used=False,
        warnings=[],
    )


def local_zip_sort_key(path: Path, roc_yyyymm: str) -> tuple[int, int, str]:
    name = path.name.lower()
    if name == f'market_{roc_yyyymm}.zip':
        rank = 0
    elif roc_yyyymm in name and '信用卡重要資訊揭露' in path.name:
        rank = 1
    elif roc_yyyymm in name:
        rank = 2
    else:
        rank = 9

    try:
        mtime = -int(path.stat().st_mtime)
    except Exception:
        mtime = 0
    return rank, mtime, name


def find_local_zip_candidates(y: int, m: int, search_dirs: list[Path]) -> list[Path]:
    roc_yyyymm = f'{y:03d}{m:02d}'
    patterns = [f'market_{roc_yyyymm}.zip', f'*{roc_yyyymm}*.zip', '*.zip']
    seen: set[Path] = set()
    out: list[Path] = []

    for base_dir in search_dirs:
        if not base_dir.exists() or not base_dir.is_dir():
            continue
        for pattern in patterns:
            for path in base_dir.glob(pattern):
                if not path.is_file():
                    continue
                if pattern == '*.zip' and roc_yyyymm not in path.name and '信用卡重要資訊揭露' not in path.name:
                    continue
                resolved = path.resolve()
                if resolved in seen:
                    continue
                seen.add(resolved)
                out.append(path)

    return sorted(out, key=lambda path: local_zip_sort_key(path, roc_yyyymm))


def fetch_market_total_by_roc(y: int, m: int, options: FetchOptions) -> dict[str, Any]:
    expected_yyyymm = roc_to_ad_yyyymm(y, m)
    if options.explicit_market and options.explicit_market['ad_yyyymm'] == expected_yyyymm:
        return options.explicit_market

    warnings: list[str] = []
    if options.prefer_local_zip:
        for path in find_local_zip_candidates(y, m, options.zip_search_dirs):
            try:
                market = parse_market_from_zip_bytes(path.read_bytes(), str(path))
            except Exception as exc:
                warnings.append(f'本機 ZIP 不可用，略過 {path}: {type(exc).__name__}: {exc}')
                continue
            if market['ad_yyyymm'] != expected_yyyymm:
                warnings.append(f'本機 ZIP 月份不符，略過 {path}: 預期 {expected_yyyymm}，實際 {market["ad_yyyymm"]}')
                continue
            return with_source_meta(
                market,
                source_kind='local_zip_auto',
                source_url=None,
                local_zip_path=str(path),
                download_mode=None,
                ssl_fallback_used=False,
                warnings=warnings,
            )

    url = build_zip_url(y, m)
    zip_bytes, meta = download_bytes(url, insecure=options.insecure)
    market = parse_market_from_zip_bytes(zip_bytes, url)
    if market['ad_yyyymm'] != expected_yyyymm:
        raise ValueError(f'銀行局月份不一致: 預期 {roc_month_text(y, m)}，實際 {market["month"]}，url={url}')

    if meta.warning:
        warnings.append(meta.warning)
    return with_source_meta(
        market,
        source_kind='download',
        source_url=url,
        local_zip_path=None,
        download_mode=meta.mode,
        ssl_fallback_used=meta.ssl_fallback_used,
        warnings=warnings,
    )


def shift_roc_month(y: int, m: int, delta: int) -> tuple[int, int]:
    total = (y + 1911) * 12 + (m - 1) + delta
    ad_y, month_index = divmod(total, 12)
    return ad_y - 1911, month_index + 1


def candidate_roc_months(max_back_months: int, base_month: str = '') -> list[tuple[int, int]]:
    if base_month:
        ad_y, month = parse_roc_month(base_month)
    else:
        today = datetime.today()
        ad_y, month = today.year - 1911, today.month
    out = []
    for offset in range(max_back_months):
        out.append(shift_roc_month(ad_y, month, -offset))
    return out


def find_latest_market_total(max_back_months: int, options: FetchOptions, base_month: str = '') -> dict[str, Any]:
    if options.explicit_market:
        return options.explicit_market

    errors = []
    for y, m in candidate_roc_months(max_back_months, base_month=base_month):
        try:
            return fetch_market_total_by_roc(y, m, options)
        except Exception as exc:
            errors.append(f'{roc_month_text(y, m)}: {type(exc).__name__}: {exc}')
    raise RuntimeError('無法找到可下載或可讀取的銀行局最新信用卡揭露資料。最近錯誤: ' + ' | '.join(errors[:8]))


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == '')


def get_main_sheet(wb):
    if MAIN_SHEET in wb.sheetnames:
        return wb[MAIN_SHEET]
    return wb[wb.sheetnames[0]]


def build_index_map(ws) -> dict[str, int]:
    normalized_headers = {normalize_header(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}
    index_map: dict[str, int] = {}
    for field, aliases in LONGFORM_FIELD_ALIASES.items():
        for alias in aliases:
            col = normalized_headers.get(normalize_header(alias))
            if col is not None:
                index_map[field] = col
                break
    required = ['yyyymm', 'ad_year', 'month_number', 'rank', 'item']
    missing = [field for field in required if field not in index_map]
    if missing:
        raise ValueError(f'{MAIN_SHEET} 缺少必要欄位: {missing}')
    return index_map


def build_month_item_rows(ws, index_map: dict[str, int]) -> dict[tuple[int, str], int]:
    mapping: dict[tuple[int, str], int] = {}
    yyyymm_col = index_map['yyyymm']
    item_col = index_map['item']
    for row_no in range(2, ws.max_row + 1):
        yyyymm = to_number(ws.cell(row_no, yyyymm_col).value)
        item = str(ws.cell(row_no, item_col).value or '').strip()
        if yyyymm is None or not item:
            continue
        mapping[(int(yyyymm), item)] = row_no
    return mapping


def set_row_value_if_present(ws, row_no: int, index_map: dict[str, int], field: str, value: Any) -> None:
    column = index_map.get(field)
    if column is not None:
        ws.cell(row_no, column).value = value


def set_month_item_metadata(ws, row_no: int, index_map: dict[str, int], *, ad_yyyymm: int, ad_year: int, month_number: int, rank: int | None, item: str) -> None:
    set_row_value_if_present(ws, row_no, index_map, 'yyyymm', ad_yyyymm)
    set_row_value_if_present(ws, row_no, index_map, 'ad_year', ad_year)
    set_row_value_if_present(ws, row_no, index_map, 'month_number', month_number)
    set_row_value_if_present(ws, row_no, index_map, 'rank', rank)
    set_row_value_if_present(ws, row_no, index_map, 'item', item)


def find_special_item_row(ws, index_map: dict[str, int], row_map: dict[tuple[int, str], int], ad_yyyymm: int, ad_year: int, month_number: int, item: str) -> int | None:
    """回傳既有月 block 中該 Item 的列號；月 block 不存在時回傳 None（不自建列）。"""
    row_no = row_map.get((ad_yyyymm, item))
    if row_no is not None:
        set_month_item_metadata(ws, row_no, index_map, ad_yyyymm=ad_yyyymm, ad_year=ad_year, month_number=month_number, rank=ITEM_RANKS.get(item), item=item)
    return row_no


def row_has_any_blank_metric(ws, row_no: int, index_map: dict[str, int], fields: list[str]) -> bool:
    for field in fields:
        column = index_map.get(field)
        if column is None:
            continue
        if is_blank(ws.cell(row_no, column).value):
            return True
    return False


def write_market_total_row(ws, row_no: int, index_map: dict[str, int], market: dict[str, Any], overwrite: bool) -> dict[str, Any]:
    written: dict[str, Any] = {}
    skipped_existing: dict[str, Any] = {}
    for field in LONGFORM_MARKET_FIELDS:
        column = index_map.get(field)
        if column is None:
            continue
        cell = ws.cell(row_no, column)
        value = market.get(field)
        if overwrite or is_blank(cell.value):
            cell.value = value
            if cell.value is not None:
                apply_default_font(cell)
            if field in PERCENT_DISPLAY_FIELDS and cell.value is not None:
                apply_percent_number_format(cell)
            written[field] = value
        else:
            skipped_existing[field] = cell.value
    return {'written': written, 'skipped_existing': skipped_existing}


def write_bank_bureau_row(ws, row_no: int, index_map: dict[str, int], market: dict[str, Any], overwrite: bool) -> dict[str, Any]:
    written: dict[str, Any] = {}
    skipped_existing: dict[str, Any] = {}
    column = index_map.get('market_total')
    if column is None:
        return {'written': written, 'skipped_existing': skipped_existing}
    cell = ws.cell(row_no, column)
    value = market.get('circulating_cards')
    if overwrite or is_blank(cell.value):
        cell.value = value
        if value is not None:
            apply_default_font(cell)
        written['market_total'] = value
    else:
        skipped_existing['market_total'] = cell.value
    return {'written': written, 'skipped_existing': skipped_existing}


def build_result(action: str, market: dict[str, Any], write_result: dict[str, Any], *, row_no: int, item: str) -> dict[str, Any]:
    return {
        'month': market['month'],
        'ad_yyyymm': market['ad_yyyymm'],
        'row': row_no,
        'item': item,
        'source_url': market.get('source_url'),
        'source_kind': market.get('source_kind'),
        'local_zip_path': market.get('local_zip_path'),
        'download_mode': market.get('download_mode'),
        'ssl_fallback_used': market.get('ssl_fallback_used', False),
        'warnings': market.get('warnings', []),
        'action': action,
        **write_result,
    }


def make_json_safe(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, list):
        return [make_json_safe(item) for item in value]
    if isinstance(value, dict):
        return {str(k): make_json_safe(v) for k, v in value.items()}
    return value


def configure_utf8_stdio() -> None:
    """盡量把 stdout / stderr 設為 UTF-8，避免中文輸出亂碼。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            if hasattr(stream, "reconfigure"):
                stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description='自動抓取銀行局信用卡重要資訊揭露資料，回補 銀行局信用卡公開資料.xlsx 空白欄位；不建立備份。')
    parser.add_argument('--workbook', default=str(DEFAULT_WORKBOOK), help='目標 Excel 檔，預設為 銀行局信用卡公開資料.xlsx')
    parser.add_argument('--lookback-months', type=int, default=24, help='自動尋找最新資料與回補舊月份時，最多往前檢查幾個月。預設 24')
    parser.add_argument('--base-month', default='', help='以指定月份為基期往前檢查空白欄位，例如 115年05月；未指定時以今天往前推 lookback months')
    parser.add_argument('--overwrite', action='store_true', help='若指定，會覆蓋既有資料；預設只補空白。')
    parser.add_argument('--insecure', action='store_true', help='直接略過 SSL 憑證驗證；若不指定，會先正常驗證，失敗時自動 fallback。')
    parser.add_argument('--zip-file', help='指定本機銀行局 ZIP 檔；若提供，優先使用此檔案。')
    parser.add_argument('--zip-search-dir', action='append', default=[], help='搜尋本機 ZIP 檔的目錄，可重複指定；預設會搜尋 . 與 input')
    parser.add_argument('--no-local-zip', action='store_true', help='不搜尋本機 ZIP，直接嘗試從網站下載。')
    parser.add_argument('--strict-workbook', action='store_true', help='若指定工作簿不可用時，直接報錯，不自動改用其他 xlsx。')
    return parser.parse_args()


def main() -> None:
    configure_utf8_stdio()
    args = parse_args()
    requested_workbook = Path(args.workbook)
    zip_search_dirs = [Path(path) for path in args.zip_search_dir] or DEFAULT_ZIP_SEARCH_DIRS

    explicit_market = None
    if args.zip_file:
        explicit_zip = Path(args.zip_file)
        if not explicit_zip.exists():
            raise FileNotFoundError(f'指定的 ZIP 檔不存在: {explicit_zip}')
        explicit_market = parse_market_from_zip_file(explicit_zip)
    else:
        explicit_zip = None

    workbook_resolution = resolve_workbook(requested_workbook, strict=args.strict_workbook)
    wb = load_workbook(workbook_resolution.path)
    ws = get_main_sheet(wb)
    index_map = build_index_map(ws)
    row_map = build_month_item_rows(ws, index_map)

    fetch_options = FetchOptions(
        insecure=args.insecure,
        prefer_local_zip=not args.no_local_zip,
        zip_search_dirs=zip_search_dirs,
        explicit_market=explicit_market,
    )

    normalized_base_month = roc_month_text(*parse_roc_month(args.base_month)) if args.base_month else ''
    latest_market = find_latest_market_total(args.lookback_months, fetch_options, base_month=normalized_base_month)

    results: list[dict[str, Any]] = []
    cache: dict[int, dict[str, Any]] = {latest_market['ad_yyyymm']: latest_market}
    if explicit_market:
        cache[explicit_market['ad_yyyymm']] = explicit_market

    candidate_months = [latest_market['ad_yyyymm']]
    seen = {latest_market['ad_yyyymm']}
    for (yyyymm, item), row_no in sorted(row_map.items(), key=lambda x: x[0][0], reverse=True):
        if item != MARKET_TOTAL_ITEM:
            continue
        if yyyymm > latest_market['ad_yyyymm'] or yyyymm in seen:
            continue
        candidate_months.append(yyyymm)
        seen.add(yyyymm)
        if len(candidate_months) >= args.lookback_months:
            break

    processed_requested_months: set[int] = set()
    processed_actual_months: set[int] = set()
    for idx, yyyymm in enumerate(candidate_months):
        if yyyymm in processed_requested_months:
            continue
        processed_requested_months.add(yyyymm)

        market = cache.get(yyyymm)
        if market is None:
            roc_y, roc_m = ad_yyyymm_to_roc(yyyymm)
            try:
                market = fetch_market_total_by_roc(roc_y, roc_m, fetch_options)
                cache[yyyymm] = market
            except Exception as exc:
                results.append({
                    'ad_yyyymm': yyyymm,
                    'action': 'skip_fetch_failed',
                    'error': f'{type(exc).__name__}: {exc}',
                })
                continue

        actual_yyyymm = int(market['ad_yyyymm'])
        if actual_yyyymm in processed_actual_months:
            results.append({
                'requested_ad_yyyymm': yyyymm,
                'ad_yyyymm': actual_yyyymm,
                'action': 'skip_duplicate_actual_month',
            })
            continue
        processed_actual_months.add(actual_yyyymm)

        market_row_no = find_special_item_row(
            ws,
            index_map,
            row_map,
            market['ad_yyyymm'],
            market['ad_year'],
            market['month_number'],
            MARKET_TOTAL_ITEM,
        )
        bank_bureau_row_no = find_special_item_row(
            ws,
            index_map,
            row_map,
            market['ad_yyyymm'],
            market['ad_year'],
            market['month_number'],
            BANK_BUREAU_ITEM,
        )

        if market_row_no is None or bank_bureau_row_no is None:
            results.append({
                'month': market['month'],
                'ad_yyyymm': market['ad_yyyymm'],
                'action': 'skip_missing_month_block',
                'warning': f'YYYYMM={market["ad_yyyymm"]} 在工作簿中沒有月 block，略過；請先由 update_credit_card_workbook.py 或 run_bank_bureau_bank_backfill.py 建立該月 13 列 block',
            })
            continue

        if idx > 0 and not args.overwrite and not row_has_any_blank_metric(ws, market_row_no, index_map, LONGFORM_MARKET_FIELDS):
            bank_bureau_column = index_map.get('market_total')
            if bank_bureau_column is not None and not is_blank(ws.cell(bank_bureau_row_no, bank_bureau_column).value):
                continue

        action = 'latest_month_update' if idx == 0 else 'backfill_blank_metrics'
        results.append(
            build_result(
                action,
                market,
                write_market_total_row(ws, market_row_no, index_map, market, overwrite=args.overwrite),
                row_no=market_row_no,
                item=MARKET_TOTAL_ITEM,
            )
        )
        results.append(
            build_result(
                action + '_bank_bureau',
                market,
                write_bank_bureau_row(ws, bank_bureau_row_no, index_map, market, overwrite=args.overwrite),
                row_no=bank_bureau_row_no,
                item=BANK_BUREAU_ITEM,
            )
        )

    if hasattr(wb, 'calculation'):
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            pass

    wb.save(workbook_resolution.path)

    payload = {
        'status': 'success',
        'requested_workbook': str(requested_workbook),
        'selected_workbook': str(workbook_resolution.path),
        'workbook_fallback_used': workbook_resolution.fallback_used,
        'backup': None,
        'backup_enabled': False,
        'latest_market_month': latest_market['month'],
        'base_month': normalized_base_month or None,
        'latest_market_source': latest_market.get('source_url') or latest_market.get('local_zip_path'),
        'latest_market_source_kind': latest_market.get('source_kind'),
        'overwrite': args.overwrite,
        'sheet': ws.title,
        'metric_labels': METRIC_LABELS,
        'zip_file': str(explicit_zip) if explicit_zip else None,
        'zip_search_dirs': [str(path) for path in zip_search_dirs],
        'summary_warnings': workbook_resolution.warnings + latest_market.get('warnings', []),
        'results': results,
    }
    print(json.dumps(make_json_safe(payload), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        print(json.dumps({
            'status': 'error',
            'error_type': type(exc).__name__,
            'error': str(exc),
        }, ensure_ascii=False, indent=2), file=sys.stderr)
        raise
