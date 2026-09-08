from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# 避免在唯讀執行沙箱中產生 __pycache__/*.pyc（會觸發 Refusing to overwrite）
sys.dont_write_bytecode = True

from bank_aliases import BANK_NAMES, BANK_ORDER, BLOCK_ITEMS, BLOCK_SIZE, MARKET_TOTAL_ITEM, SPECIAL_ITEMS, year_block_key
from percent_utils import PERCENT_DECIMAL_FIELDS, apply_percent_number_format, assert_percent_sane, normalize_percent_value
from workbook_block_helpers import (
    append_year_block as shared_append_year_block,
    apply_default_font,
    audit_block_percent_ratios,
    canonical_block_item,
    copy_formula_cells,
    find_block,
    find_month_block as shared_find_month_block,
    find_or_create_month_block as shared_find_or_create_month_block,
    normalize_item_labels,
    refresh_block_metadata as shared_refresh_block_metadata,
    validate_block_items,
)

WORKBOOK_DEFAULT = "銀行局信用卡公開資料.xlsx"
MAIN_SHEET = "歷史資料(年+月)"

FIELD_ALIASES = {
    "yyyymm": ["YYYYMM"],
    "ad_year": ["年度"],
    "month_number": ["月份"],
    "rank": ["Rank", "排序編號"],
    "bank": ["Bank"],
    "item": ["Item"],
    "circulating_cards": ["流通卡數"],
    "valid_cards": ["有效卡數"],
    "new_cards_this_month": ["當月發卡數"],
    "cancelled_cards_this_month": ["當月停卡數"],
    "revolving_balance_million": ["循環信用餘額"],
    "installment_balance_not_yet_due_million": ["未到期分期付款餘額"],
    "signed_amount_million": ["當月簽帳金額"],
    "cash_advance_amount_million": ["當月預借現金金額"],
    "overdue_3m_ratio_percent": [
        "逾期三個月以上比率",
        "逾期三個月比率",
        "逾期三個月以上帳款占應收帳款餘額（含催收款）之比率(%)",
        "逾期三個月以上帳款占應收帳款餘額(含催收款)之比率(%)",
        "逾期三個月以上帳款佔應收帳款餘額（含催收款）之比率(%)",
        "逾期三個月以上帳款佔應收帳款餘額(含催收款)之比率(%)",
    ],
    "overdue_6m_ratio_percent": [
        "逾期六個月以上比率",
        "逾期六個月比率",
        "逾期六個月以上帳款占應收帳款餘額（含催收款）之比率(%)",
        "逾期六個月以上帳款占應收帳款餘額(含催收款)之比率(%)",
        "逾期六個月以上帳款佔應收帳款餘額（含催收款）之比率(%)",
        "逾期六個月以上帳款佔應收帳款餘額(含催收款)之比率(%)",
    ],
    "allowance_coverage_ratio_percent": ["備抵呆帳提足率", "備抵呆帳提足率(%)"],
    "charge_off_amount_this_month_million": ["當月轉銷呆帳金額", "當月轉銷呆帳金額"],
    "charge_off_amount_ytd_million": ["當年度累計轉銷呆帳金額", "當年度累計轉銷呆帳金額"],
    "debit_card_signed_amount_million": ["當月轉帳卡簽帳金額", "轉帳卡簽帳金額"],
    "avg_cards_per_person": ["平均每人持卡張數"],
    "market_total": ["市場總計"],
    "top5_circulating_cards": ["TOP 5 流通卡"],
    "top10_circulating_cards": ["TOP 10 流通卡"],
    "top5_valid_cards": ["TOP 5 有效卡"],
    "top10_valid_cards": ["TOP 10 有效卡"],
    "top5_new_cards": ["TOP 5 當月發卡數"],
    "top10_new_cards": ["TOP 10 當月發卡數"],
    "top5_cancelled_cards": ["TOP 5 當月停卡數"],
    "top10_cancelled_cards": ["TOP 10 當月停卡數"],
    "top5_revolving_balance": ["TOP 5 循環信用餘額"],
    "top10_revolving_balance": ["TOP 10 循環信用餘額"],
    "top5_signed_amount": ["TOP 5 當月簽帳金額"],
    "top10_signed_amount": ["TOP 10 當月簽帳金額"],
}

PERCENT_DISPLAY_FIELDS = PERCENT_DECIMAL_FIELDS
INTEGER_DISPLAY_FIELDS = {
    "installment_balance_not_yet_due_million",
    "cash_advance_amount_million",
    "charge_off_amount_this_month_million",
    "charge_off_amount_ytd_million",
    "debit_card_signed_amount_million",
}
INTEGER_NUMBER_FORMAT = '#,##0'


def apply_integer_number_format(cell) -> None:
    cell.number_format = INTEGER_NUMBER_FORMAT


WRITABLE_FIELDS = [
    "circulating_cards",
    "valid_cards",
    "new_cards_this_month",
    "cancelled_cards_this_month",
    "revolving_balance_million",
    "installment_balance_not_yet_due_million",
    "signed_amount_million",
    "cash_advance_amount_million",
    "overdue_3m_ratio_percent",
    "overdue_6m_ratio_percent",
    "allowance_coverage_ratio_percent",
    "charge_off_amount_this_month_million",
    "charge_off_amount_ytd_million",
    "debit_card_signed_amount_million",
]

SUMMARY_REQUIRED_FIELDS = [
    "circulating_cards",
    "valid_cards",
    "new_cards_this_month",
    "cancelled_cards_this_month",
    "signed_amount_million",
    "revolving_balance_million",
]

ANNUAL_END_VALUE_FIELDS = [
    "circulating_cards",
    "valid_cards",
    "revolving_balance_million",
    "installment_balance_not_yet_due_million",
    "overdue_3m_ratio_percent",
    "overdue_6m_ratio_percent",
    "allowance_coverage_ratio_percent",
    "charge_off_amount_ytd_million",
]

ANNUAL_SUM_FIELDS = [
    "new_cards_this_month",
    "cancelled_cards_this_month",
    "signed_amount_million",
    "cash_advance_amount_million",
    "charge_off_amount_this_month_million",
    "debit_card_signed_amount_million",
]

ANNUAL_REQUIRED_MONTHLY_ITEMS = BLOCK_ITEMS

ANNUAL_REQUIRED_MONTHLY_FIELDS = {
    **{BANK_NAMES[key]: SUMMARY_REQUIRED_FIELDS for key in BANK_ORDER},
    MARKET_TOTAL_ITEM: SUMMARY_REQUIRED_FIELDS + ["avg_cards_per_person"],
}

# 市場總計列上的公式欄：市場總計（SUMIFS）與 TOP5/TOP10 占比；值為 (分子指標欄, top_n)，供無公式樣板時改算數值。
MARKET_ROW_FORMULA_FIELDS = {
    "market_total": ("circulating_cards", 0),
    "top5_circulating_cards": ("circulating_cards", 5),
    "top10_circulating_cards": ("circulating_cards", 10),
    "top5_valid_cards": ("valid_cards", 5),
    "top10_valid_cards": ("valid_cards", 10),
    "top5_new_cards": ("new_cards_this_month", 5),
    "top10_new_cards": ("new_cards_this_month", 10),
    "top5_cancelled_cards": ("cancelled_cards_this_month", 5),
    "top10_cancelled_cards": ("cancelled_cards_this_month", 10),
    "top5_revolving_balance": ("revolving_balance_million", 5),
    "top10_revolving_balance": ("revolving_balance_million", 10),
    "top5_signed_amount": ("signed_amount_million", 5),
    "top10_signed_amount": ("signed_amount_million", 10),
}

def parse_roc_month(month_text: str) -> tuple[int, int]:
    match = re.search(r"(\d{3})\D*(\d{1,2})", str(month_text))
    if not match:
        raise ValueError(f"無法解析月份：{month_text}")
    year, month = int(match.group(1)), int(match.group(2))
    if not 1 <= month <= 12:
        raise ValueError(f"月份不合法：{month_text}")
    return year, month


def roc_month_text(year: int, month: int) -> str:
    return f"{year}年{month:02d}月"


def ad_year(roc_year: int) -> int:
    return roc_year + 1911


def roc_to_ad_yyyymm(roc_year: int, month: int) -> int:
    return ad_year(roc_year) * 100 + month


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").replace("\u3000", "")).strip()


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def first_present(payload: dict[str, Any], keys: list[str], default: Any = None) -> Any:
    for key in keys:
        if key in payload and payload[key] is not None:
            return payload[key]
    return default


def as_number(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    text = str(value).strip().replace(",", "")
    if text in ("", "-", "—", "－"):
        return None
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except Exception:
        return value


def derive_million_from_summary(rec: dict[str, Any], million_keys: list[str], thousand_keys: list[str]) -> Any:
    million = as_number(first_present(rec, million_keys))
    if million is not None:
        return million
    thousand = as_number(first_present(rec, thousand_keys))
    if isinstance(thousand, (int, float)):
        return thousand / 1000
    return None


def decode_subprocess_output(data: bytes | None) -> str:
    if not data:
        return ""
    for encoding in ("utf-8", "utf-8-sig", "cp950", "big5"):
        try:
            return data.decode(encoding)
        except Exception:
            pass
    return data.decode("utf-8", errors="replace")


def configure_utf8_stdio() -> None:
    """盡量把 stdout / stderr 設為 UTF-8，避免中文輸出亂碼。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            if hasattr(stream, "reconfigure"):
                stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

def normalize_yyyymm(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        if 190001 <= value <= 299912 and 1 <= value % 100 <= 12:
            return value
        return None
    if isinstance(value, float) and value.is_integer():
        return normalize_yyyymm(int(value))

    text = str(value).strip()
    if re.fullmatch(r"20\d{4}", text):
        yyyymm = int(text)
        if 1 <= yyyymm % 100 <= 12:
            return yyyymm
        return None

    try:
        roc_year, month = parse_roc_month(text)
        return roc_to_ad_yyyymm(roc_year, month)
    except Exception:
        return None


def header_index_map(ws, header_aliases: dict[str, list[str]]) -> dict[str, int]:
    alias_lookup: dict[str, str] = {}
    for canonical, aliases in header_aliases.items():
        for alias in aliases:
            alias_lookup[normalize_header(alias)] = canonical

    out: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        normalized = normalize_header(ws.cell(1, column).value)
        canonical = alias_lookup.get(normalized)
        if canonical:
            out[canonical] = column
    return out


def require_headers(index_map: dict[str, int], required: list[str], sheet_name: str) -> None:
    missing = [field for field in required if field not in index_map]
    if missing:
        raise ValueError(f"{sheet_name} 缺少欄位：{missing}")


def get_main_sheet(wb):
    if MAIN_SHEET not in wb.sheetnames:
        raise ValueError(f"工作簿缺少工作表：{MAIN_SHEET}；現有工作表：{wb.sheetnames}")
    return wb[MAIN_SHEET]


def normalize_json_records(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    if not isinstance(payload, dict):
        raise ValueError("JSON 格式不支援：最外層應為 list 或 dict")

    for key in ["results", "banks", "data", "items", "summary"]:
        value = payload.get(key)
        if isinstance(value, list):
            return [item for item in value if isinstance(item, dict)]
        if isinstance(value, dict):
            records = []
            for bank_key, item in value.items():
                if isinstance(item, dict):
                    record = dict(item)
                    record.setdefault("bank_key", bank_key)
                    records.append(record)
            return records

    records = []
    for bank_key, item in payload.items():
        if isinstance(item, dict):
            record = dict(item)
            record.setdefault("bank_key", bank_key)
            records.append(record)
    if records:
        return records
    raise ValueError("JSON 找不到可解析的銀行資料清單")


def merged_record(rec: dict[str, Any]) -> dict[str, Any]:
    merged = dict(rec)
    metrics = rec.get("metrics")
    if isinstance(metrics, dict):
        for key, value in metrics.items():
            merged.setdefault(key, value)
    return merged


def load_summary_results(summary_path: Path, allow_partial: bool = False) -> dict[str, dict[str, Any]]:
    if summary_path.suffix.lower() != ".json":
        raise ValueError(f"--summary 必須是 JSON 檔：{summary_path}")
    with summary_path.open("r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)

    records = normalize_json_records(payload)
    name_to_key = {value: key for key, value in BANK_NAMES.items()}
    out: dict[str, dict[str, Any]] = {}

    for raw_rec in records:
        rec = merged_record(raw_rec)
        raw_key = first_present(rec, ["bank_key", "bank_code", "bank", "key", "code"])
        bank_name = first_present(rec, ["bank_name", "銀行名稱", "name"])
        key = str(raw_key).strip().lower() if raw_key else ""
        if key not in BANK_ORDER and bank_name in name_to_key:
            key = name_to_key[str(bank_name)]
        if key not in BANK_ORDER:
            continue

        signed_million = derive_million_from_summary(rec, ["signed_amount_million", "當月簽帳金額(百萬元)"], ["signed_amount_thousand", "當月簽帳金額(千元)", "signed_amount"])
        revolving_million = derive_million_from_summary(rec, ["revolving_balance_million", "循環信用餘額(百萬元)"], ["revolving_balance_thousand", "循環信用餘額(千元)", "revolving_balance"])
        debit_million = derive_million_from_summary(rec, ["debit_card_signed_amount_million", "當月轉帳卡簽帳金額(百萬)"], ["debit_card_signed_amount_thousand", "轉帳卡簽帳金額(千元)"])

        allowance_coverage = as_number(first_present(rec, ["allowance_coverage_ratio_percent", "備抵呆帳提足率"]))
        if allowance_coverage is None:
            allowance_coverage = normalize_percent_value(first_present(rec, ["allowance_coverage_ratio_percent", "備抵呆帳提足率"]))

        def overdue_ratio_value(field: str, aliases: list[str]) -> Any:
            raw = first_present(rec, [field, *aliases])
            for unit_key in ("metrics_units", "metric_units"):
                units = rec.get(unit_key)
                if isinstance(units, dict) and units.get(field) == "decimal_ratio":
                    # run_all_banks 標準化後的 summary：值已是小數比率，直接採用
                    return as_number(raw)
            # 未標準化來源（手工 JSON、中文欄位）：值為百分比顯示數字（0.12 代表 0.12%），一律 /100
            return normalize_percent_value(raw, force_percent_input=True)

        out[key] = {
            "bank_key": key,
            "bank_name": bank_name or BANK_NAMES[key],
            "status": first_present(rec, ["status", "執行結果"], "success"),
            "data_month": first_present(rec, ["data_month", "資料年月"]),
            "base_date": first_present(rec, ["base_date", "基準日"]),
            "circulating_cards": as_number(first_present(rec, ["circulating_cards", "流通卡數"])),
            "valid_cards": as_number(first_present(rec, ["valid_cards", "有效卡數"])),
            "new_cards_this_month": as_number(first_present(rec, ["new_cards_this_month", "當月發卡數"])),
            "cancelled_cards_this_month": as_number(first_present(rec, ["cancelled_cards_this_month", "當月停卡數"])),
            "signed_amount_million": signed_million,
            "revolving_balance_million": revolving_million,
            "installment_balance_not_yet_due_million": derive_million_from_summary(rec, ["installment_balance_not_yet_due_million", "未到期分期付款餘額(百萬元)"], ["installment_balance_not_yet_due_thousand", "未到期分期付款餘額(千元)", "未到期分期付款餘額"]),
            "cash_advance_amount_million": derive_million_from_summary(rec, ["cash_advance_amount_million", "當月預借現金金額(百萬元)"], ["cash_advance_amount_thousand", "當月預借現金金額(千元)", "cash_advance_amount"]),
            "overdue_3m_ratio_percent": overdue_ratio_value("overdue_3m_ratio_percent", ["逾期三個月以上比率", "逾期三個月比率"]),
            "overdue_6m_ratio_percent": overdue_ratio_value("overdue_6m_ratio_percent", ["逾期六個月以上比率", "逾期六個月比率"]),
            "allowance_coverage_ratio_percent": allowance_coverage,
            "charge_off_amount_this_month_million": derive_million_from_summary(rec, ["charge_off_amount_this_month_million", "當月轉銷呆帳金額(百萬元)"], ["charge_off_amount_this_month_thousand", "當月轉銷呆帳金額(千元)", "當月轉銷呆帳金額"]),
            "charge_off_amount_ytd_million": derive_million_from_summary(rec, ["charge_off_amount_ytd_million", "當年度累計轉銷呆帳金額(百萬元)"], ["charge_off_amount_ytd_thousand", "當年度累計轉銷呆帳金額(千元)", "當年度累計轉銷呆帳金額"]),
            "debit_card_signed_amount_million": debit_million,
            "notes": first_present(rec, ["notes", "failure_reason", "失敗原因"]),
            "source_type": first_present(rec, ["source_type", "來源類型"]),
            "source_url": first_present(rec, ["source_url", "來源網址"]),
            "error_count": as_number(first_present(rec, ["error_count", "錯誤數"], 0)),
        }

    missing = [key for key in BANK_ORDER if key not in out]
    if missing and not allow_partial:
        raise ValueError(f"summary JSON 缺少銀行資料：{missing}")
    return out


def enrich_ctbc_debit_card_amount(results: dict[str, dict[str, Any]]) -> None:
    item = results.get("ctbc")
    if not item or item.get("debit_card_signed_amount_million") is not None:
        return

    candidates = []
    if item.get("source_url"):
        candidates.append(Path(str(item["source_url"])))
    candidates.append(Path("中信信用卡資料.xlsx"))
    src = next((path for path in candidates if path.exists()), None)
    if not src:
        return

    wb = load_workbook(src, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            text = " ".join("" if value is None else str(value).strip() for value in row)
            if "轉帳卡簽帳金額" not in text:
                continue
            numbers = [value for value in row if isinstance(value, (int, float))]
            if numbers:
                item["debit_card_signed_amount_million"] = numbers[-1] / 1000
                return
    finally:
        wb.close()


def infer_target_month_from_results(results: dict[str, dict[str, Any]]) -> str:
    months: set[tuple[int, int]] = set()
    for item in results.values():
        data_month = item.get("data_month")
        if not data_month:
            continue
        months.add(parse_roc_month(str(data_month)))
    if not months:
        raise ValueError("summary JSON 中找不到可推定 target month 的 data_month")
    # 各家銀行公布時間不同，取最新月為 target，落後月份留給 backfill。
    year, month = max(months)
    return roc_month_text(year, month)


def determine_backfill_base_month(target_month: str | None, results: dict[str, dict[str, Any]] | None = None) -> str | None:
    if target_month:
        return target_month
    if results:
        return infer_target_month_from_results(results)
    return None


def find_month_block(ws, index_map: dict[str, int], yyyymm: int) -> dict[str, Any] | None:
    return shared_find_month_block(ws, index_map, yyyymm, normalize_yyyymm=normalize_yyyymm)


def find_or_create_month_block(ws, index_map: dict[str, int], ad_year: int, month_number: int) -> dict[str, Any]:
    return shared_find_or_create_month_block(ws, index_map, ad_year=ad_year, month_number=month_number, normalize_yyyymm=normalize_yyyymm)


def refresh_block_metadata(ws, index_map: dict[str, int], block_info: dict[str, Any], ad_year: int, month_number: int) -> None:
    shared_refresh_block_metadata(ws, index_map, block_info, ad_year=ad_year, month_number=month_number)


def collect_month_row_groups(ws, index_map: dict[str, int]) -> dict[int, list[int]]:
    groups: dict[int, list[int]] = {}
    for row_no in range(2, ws.max_row + 1):
        yyyymm = normalize_yyyymm(ws.cell(row_no, index_map["yyyymm"]).value)
        if yyyymm is not None:
            groups.setdefault(yyyymm, []).append(row_no)
    return groups


def classify_month_rows(ws, index_map: dict[str, int], yyyymm: int, rows: list[int]) -> str:
    try:
        validate_block_items(ws, index_map, rows, yyyymm)
        return "full_block"
    except ValueError:
        pass
    if len(rows) < BLOCK_SIZE:
        items = [canonical_block_item(ws.cell(row_no, index_map["item"]).value) for row_no in rows]
        if len(set(items)) == len(items) and all(item in BLOCK_ITEMS for item in items):
            return "partial_subset"
    return "needs_manual_review"


def detect_partial_month_blocks(ws, index_map: dict[str, int]) -> list[int]:
    groups = collect_month_row_groups(ws, index_map)
    return sorted(
        yyyymm for yyyymm, rows in groups.items()
        if classify_month_rows(ws, index_map, yyyymm, rows) != "full_block"
    )


def repair_partial_blocks(ws, index_map: dict[str, int]) -> dict[str, Any]:
    groups = collect_month_row_groups(ws, index_map)
    classified = {yyyymm: classify_month_rows(ws, index_map, yyyymm, rows) for yyyymm, rows in groups.items()}

    # openpyxl delete_rows 不會平移下方既有公式的參照，因此只允許刪除
    # 位於「所有受保護列」之後（工作表尾端）的 partial block。
    # 受保護列 = 完整 block、需人工確認的列，以及月分組以外仍有內容的列（例如年度 block）。
    protected_max_row = 1
    month_rows = {row for rows in groups.values() for row in rows}
    for yyyymm, rows in groups.items():
        if classified[yyyymm] != "partial_subset":
            protected_max_row = max(protected_max_row, max(rows))
    for row_no in range(2, ws.max_row + 1):
        if row_no in month_rows:
            continue
        if not is_blank(ws.cell(row_no, index_map["yyyymm"]).value) or not is_blank(ws.cell(row_no, index_map["item"]).value):
            protected_max_row = max(protected_max_row, row_no)

    deleted: list[dict[str, Any]] = []
    needs_manual_review: list[dict[str, Any]] = []
    rows_to_delete: list[int] = []

    for yyyymm in sorted(groups):
        status = classified[yyyymm]
        if status == "full_block":
            continue
        rows = groups[yyyymm]
        if status == "needs_manual_review":
            needs_manual_review.append({"yyyymm": yyyymm, "rows": rows, "reason": "not_partial_subset"})
            continue
        if min(rows) <= protected_max_row:
            needs_manual_review.append({"yyyymm": yyyymm, "rows": rows, "reason": "not_at_tail_refuse_delete"})
            continue
        removed: list[dict[str, Any]] = []
        for row_no in rows:
            values: dict[str, Any] = {}
            for field, column in index_map.items():
                if field in ("yyyymm", "ad_year", "month_number", "rank", "item"):
                    continue
                value = ws.cell(row_no, column).value
                if not is_blank(value):
                    values[FIELD_ALIASES[field][0]] = value
            removed.append({
                "row": row_no,
                "item": str(ws.cell(row_no, index_map["item"]).value or ""),
                "values": values,
            })
        deleted.append({"yyyymm": yyyymm, "rows": rows, "removed": removed})
        rows_to_delete.extend(rows)

    for row_no in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(row_no, 1)

    return {
        "full_block_count": sum(1 for status in classified.values() if status == "full_block"),
        "deleted": deleted,
        "needs_manual_review": needs_manual_review,
        "deleted_row_count": len(rows_to_delete),
    }


def validate_results_for_write(results: dict[str, dict[str, Any]], allow_partial: bool = False) -> None:
    missing_banks = [key for key in BANK_ORDER if key not in results]
    if missing_banks and not allow_partial:
        raise ValueError(f"summary JSON 缺少銀行資料：{missing_banks}")


def write_bank_rows_to_block(ws, index_map: dict[str, int], block_info: dict[str, Any], results: dict[str, dict[str, Any]], target_month: str | None = None) -> dict[str, Any]:
    updated_rows = 0
    written_fields: set[str] = set()
    skipped_banks: list[str] = []

    target_ym = parse_roc_month(target_month) if target_month else None

    for bank_key in BANK_ORDER:
        item = results.get(bank_key)
        if item is None:
            continue
        # 只把「與目標月份相符」的銀行寫進這個 block。落後月份的銀行資料
        # 若貼進最新月 block 會造成資料錯置，先跳過並回報，交給 backfill 補正確月份。
        if target_ym is not None:
            data_month = item.get("data_month")
            if data_month and parse_roc_month(str(data_month)) != target_ym:
                skipped_banks.append(f"{BANK_NAMES[bank_key]}({data_month})")
                continue
        row_no = block_info["item_to_row"][BANK_NAMES[bank_key]]
        for field in WRITABLE_FIELDS:
            column = index_map.get(field)
            if column is None:
                continue
            assert_percent_sane(field, item.get(field), context=f"{target_month or ''} {BANK_NAMES[bank_key]}".strip())
            cell = ws.cell(row_no, column)
            cell.value = item.get(field)
            if cell.value is not None:
                apply_default_font(cell)
            if field in PERCENT_DISPLAY_FIELDS and cell.value is not None:
                apply_percent_number_format(cell)
            elif field in INTEGER_DISPLAY_FIELDS and cell.value is not None:
                apply_integer_number_format(cell)
            written_fields.add(field)
        updated_rows += 1

    return {
        "bank_rows_updated": updated_rows,
        "written_fields": [FIELD_ALIASES[field][0] for field in WRITABLE_FIELDS if field in written_fields],
        "skipped_banks": skipped_banks,
    }


def normalize_year_value(value: Any) -> int | None:
    """年度 block 的 YYYYMM 欄：'2025--'（現行寫法）或 2025 / '2025'（舊寫法）。"""
    if value is None:
        return None
    if isinstance(value, int):
        return value if 1900 <= value <= 2999 else None
    if isinstance(value, float) and value.is_integer():
        return normalize_year_value(int(value))
    m = re.fullmatch(r"(20\d{2})(?:--)?", str(value).strip())
    return int(m.group(1)) if m else None


def find_year_block(ws, index_map: dict[str, int], ad_year_value: int) -> dict[str, Any] | None:
    return find_block(ws, index_map, year_block_key(ad_year_value), match=lambda v: normalize_year_value(v) == ad_year_value, block_label="年度 block ")


def append_year_block(ws, index_map: dict[str, int], ad_year_value: int) -> dict[str, Any]:
    return shared_append_year_block(ws, index_map, ad_year=ad_year_value)


def find_or_create_year_block(ws, index_map: dict[str, int], ad_year_value: int) -> dict[str, Any]:
    found = find_year_block(ws, index_map, ad_year_value)
    if found:
        return found
    return append_year_block(ws, index_map, ad_year_value)


def refresh_year_block_metadata(ws, index_map: dict[str, int], block_info: dict[str, Any], ad_year_value: int) -> None:
    shared_refresh_block_metadata(ws, index_map, block_info, ad_year=ad_year_value, month_number="--")


def monthly_blocks_by_year(ws, index_map: dict[str, int], ad_year_value: int) -> dict[int, dict[str, Any]]:
    out: dict[int, dict[str, Any]] = {}
    for month in range(1, 13):
        block = find_month_block(ws, index_map, ad_year_value * 100 + month)
        if block is not None:
            out[month] = block
    return out


def get_cell_numeric(ws, row_no: int, index_map: dict[str, int], field: str) -> float | int | None:
    column = index_map.get(field)
    if column is None:
        return None
    return as_number(ws.cell(row_no, column).value)


def year_completeness_report(ws, index_map: dict[str, int], ad_year_value: int) -> dict[str, Any]:
    # 年度完整性：同一年 1–12 月 block 都存在，且年度整理需要的月資料欄位不可為空。
    month_blocks = monthly_blocks_by_year(ws, index_map, ad_year_value)
    missing_months = [month for month in range(1, 13) if month not in month_blocks]
    missing_details: list[dict[str, Any]] = []

    for month in range(1, 13):
        block = month_blocks.get(month)
        if block is None:
            continue
        for item in ANNUAL_REQUIRED_MONTHLY_ITEMS:
            row_no = block["item_to_row"].get(item)
            if row_no is None:
                missing_details.append({"month": month, "item": item, "reason": "missing_item_row"})
                continue
            for field in ANNUAL_REQUIRED_MONTHLY_FIELDS.get(item, []):
                column = index_map.get(field)
                if column is None:
                    continue
                if is_blank(ws.cell(row_no, column).value):
                    missing_details.append({"month": month, "item": item, "field": field, "reason": "blank_required_field"})

    return {
        "year": ad_year_value,
        "complete": not missing_months and not missing_details,
        "missing_months": missing_months,
        "missing_details": missing_details,
        "month_blocks": month_blocks,
    }


def december_field_value(ws, month_blocks: dict[int, dict[str, Any]], index_map: dict[str, int], item_name: str, field: str) -> Any:
    december = month_blocks.get(12)
    if december is None:
        return None
    row_no = december["item_to_row"][item_name]
    column = index_map.get(field)
    if column is None:
        return None
    return ws.cell(row_no, column).value


def sum_monthly_field(ws, month_blocks: dict[int, dict[str, Any]], index_map: dict[str, int], item_name: str, field: str) -> float | int | None:
    total = 0.0
    found = False
    for month in range(1, 13):
        block = month_blocks.get(month)
        if block is None:
            continue
        row_no = block["item_to_row"][item_name]
        value = get_cell_numeric(ws, row_no, index_map, field)
        if isinstance(value, (int, float)):
            total += value
            found = True
    if not found:
        return None
    return int(total) if total.is_integer() else total


def build_annual_record_for_item(ws, month_blocks: dict[int, dict[str, Any]], index_map: dict[str, int], item_name: str) -> dict[str, Any]:
    record: dict[str, Any] = {}
    for field in ANNUAL_END_VALUE_FIELDS:
        if field in index_map:
            record[field] = december_field_value(ws, month_blocks, index_map, item_name, field)
    for field in ANNUAL_SUM_FIELDS:
        if field in index_map:
            record[field] = sum_monthly_field(ws, month_blocks, index_map, item_name, field)
    return record


def find_annual_formula_source_row(ws, index_map: dict[str, int], target_year: int) -> int | None:
    """找最接近年度的其他年度 block 市場總計列（帶 TOP 公式）當公式樣板。"""
    candidates: list[tuple[int, int]] = []
    for row_no in range(2, ws.max_row + 1):
        if canonical_block_item(ws.cell(row_no, index_map["item"]).value) != MARKET_TOTAL_ITEM:
            continue
        year_value = normalize_year_value(ws.cell(row_no, index_map["yyyymm"]).value)
        if year_value is None or year_value == target_year:
            continue
        if any(
            ws.cell(row_no, index_map[field]).data_type == "f" and isinstance(ws.cell(row_no, index_map[field]).value, str)
            for field in MARKET_ROW_FORMULA_FIELDS
            if field in index_map
        ):
            candidates.append((abs(year_value - target_year), row_no))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0])
    return candidates[0][1]


def annual_bank_bureau_ratio(ws, month_blocks: dict[int, dict[str, Any]], index_map: dict[str, int], field: str, top_n: int) -> float | None:
    if field in {"circulating_cards", "valid_cards", "revolving_balance_million"}:
        if 12 not in month_blocks:
            return None
        denominator = get_cell_numeric(ws, month_blocks[12]["item_to_row"][MARKET_TOTAL_ITEM], index_map, field)
        values = [get_cell_numeric(ws, month_blocks[12]["item_to_row"][BANK_NAMES[key]], index_map, field) for key in BANK_ORDER]
    elif field in {"new_cards_this_month", "cancelled_cards_this_month", "signed_amount_million"}:
        denominator = sum_monthly_field(ws, month_blocks, index_map, MARKET_TOTAL_ITEM, field)
        values = [sum_monthly_field(ws, month_blocks, index_map, BANK_NAMES[key], field) for key in BANK_ORDER]
    else:
        return None
    if not isinstance(denominator, (int, float)) or denominator == 0:
        return None
    usable = sorted([float(v) for v in values if isinstance(v, (int, float))], reverse=True)
    return sum(usable[:top_n]) / float(denominator)


def repair_annual_market_row_formulas(ws, index_map: dict[str, int], block_info: dict[str, Any]) -> dict[str, Any]:
    target_row = block_info["item_to_row"][MARKET_TOTAL_ITEM]
    source_row = find_annual_formula_source_row(ws, index_map, normalize_year_value(block_info["yyyymm"]))
    if source_row is None:
        return {"formula_source_row": None, "formulas_written": []}
    columns = {index_map[field]: field for field in MARKET_ROW_FORMULA_FIELDS if field in index_map}
    written, _ = copy_formula_cells(ws, source_row, target_row, list(columns), overwrite=True)
    return {"formula_source_row": source_row, "formulas_written": [columns[c] for c in written]}


def write_annual_rows_to_block(ws, index_map: dict[str, int], block_info: dict[str, Any], month_blocks: dict[int, dict[str, Any]]) -> dict[str, Any]:
    written_items: dict[str, list[str]] = {}
    for bank_key in BANK_ORDER:
        item_name = BANK_NAMES[bank_key]
        row_no = block_info["item_to_row"][item_name]
        record = build_annual_record_for_item(ws, month_blocks, index_map, item_name)
        fields_written: list[str] = []
        for field, value in record.items():
            column = index_map.get(field)
            if column is None:
                continue
            cell = ws.cell(row_no, column)
            cell.value = value
            if cell.value is not None:
                apply_default_font(cell)
            if field in PERCENT_DISPLAY_FIELDS and cell.value is not None:
                apply_percent_number_format(cell)
            elif field in INTEGER_DISPLAY_FIELDS and cell.value is not None:
                apply_integer_number_format(cell)
            fields_written.append(field)
        written_items[item_name] = fields_written

    market_row = block_info["item_to_row"][MARKET_TOTAL_ITEM]
    market_record = build_annual_record_for_item(ws, month_blocks, index_map, MARKET_TOTAL_ITEM)
    market_written: list[str] = []
    for field, value in market_record.items():
        column = index_map.get(field)
        if column is None:
            continue
        cell = ws.cell(market_row, column)
        cell.value = value
        if cell.value is not None:
            apply_default_font(cell)
        if field in PERCENT_DISPLAY_FIELDS and cell.value is not None:
            apply_percent_number_format(cell)
        elif field in INTEGER_DISPLAY_FIELDS and cell.value is not None:
            apply_integer_number_format(cell)
        market_written.append(field)
    if "avg_cards_per_person" in index_map:
        value = december_field_value(ws, month_blocks, index_map, MARKET_TOTAL_ITEM, "avg_cards_per_person")
        avg_cell = ws.cell(market_row, index_map["avg_cards_per_person"])
        avg_cell.value = value
        if value is not None:
            apply_default_font(avg_cell)
        market_written.append("avg_cards_per_person")

    # 市場總計 / TOP5 / TOP10 欄優先複製其他年度 block 的公式；沒有樣板時改寫計算值。
    formula_result = repair_annual_market_row_formulas(ws, index_map, block_info)
    if formula_result["formulas_written"]:
        market_written.extend(formula_result["formulas_written"])
    else:
        for formula_field, (metric_field, top_n) in MARKET_ROW_FORMULA_FIELDS.items():
            column = index_map.get(formula_field)
            if column is None:
                continue
            cell = ws.cell(market_row, column)
            if top_n == 0:
                cell.value = december_field_value(ws, month_blocks, index_map, MARKET_TOTAL_ITEM, metric_field)
            else:
                cell.value = annual_bank_bureau_ratio(ws, month_blocks, index_map, metric_field, top_n)
                apply_percent_number_format(cell)
            apply_default_font(cell)
            market_written.append(formula_field)
    written_items[MARKET_TOTAL_ITEM] = market_written

    return {"written_items": written_items, "market_row_formula_result": formula_result}


def auto_sync_annual_if_ready(ws, index_map: dict[str, int], ad_year_value: int) -> dict[str, Any]:
    report = year_completeness_report(ws, index_map, ad_year_value)
    if not report["complete"]:
        return {
            "triggered": False,
            "year": ad_year_value,
            "missing_months": report["missing_months"],
            "missing_details": report["missing_details"],
        }
    block_info = find_or_create_year_block(ws, index_map, ad_year_value)
    refresh_year_block_metadata(ws, index_map, block_info, ad_year_value)
    write_summary = write_annual_rows_to_block(ws, index_map, block_info, report["month_blocks"])
    return {
        "triggered": True,
        "year": ad_year_value,
        "block_created": block_info["created"],
        "block_start_row": block_info["start_row"],
        **write_summary,
    }


def audit_percent_ratios(ws, index_map: dict[str, int]) -> dict[str, Any]:
    """掃描所有完整月 block 的百分比欄（絕對範圍＋市場總計 vs 十家銀行相對檢查），回報異常格。"""
    anomalies: list[dict[str, Any]] = []
    audited = 0
    for yyyymm, rows in sorted(collect_month_row_groups(ws, index_map).items()):
        if classify_month_rows(ws, index_map, yyyymm, rows) != "full_block":
            continue
        block_info = find_month_block(ws, index_map, yyyymm)
        if block_info is None:
            continue
        audited += 1
        anomalies.extend(audit_block_percent_ratios(ws, index_map, block_info, sorted(PERCENT_DECIMAL_FIELDS)))
    return {"blocks_audited": audited, "anomaly_count": len(anomalies), "anomalies": anomalies}


def verify_written_block(workbook_path: Path, target_month: str, results: dict[str, dict[str, Any]] | None, skipped_banks: list[str]) -> dict[str, Any]:
    # 存檔後重新載入工作簿做讀回驗證，避免只憑 in-memory 狀態與委派腳本 stdout 判定成功。
    roc_year, month_number = parse_roc_month(target_month)
    yyyymm = roc_to_ad_yyyymm(roc_year, month_number)
    wb = load_workbook(workbook_path)
    try:
        ws = get_main_sheet(wb)
        index_map = header_index_map(ws, FIELD_ALIASES)
        try:
            block_info = find_month_block(ws, index_map, yyyymm)
        except ValueError as exc:
            return {"status": "failed", "yyyymm": yyyymm, "reason": str(exc)}
        if block_info is None:
            return {"status": "failed", "yyyymm": yyyymm, "reason": f"找不到 YYYYMM={yyyymm} 的月 block"}

        skipped_names = {name.split("(")[0] for name in skipped_banks}
        mismatches: list[dict[str, Any]] = []
        verified_banks: list[str] = []
        for bank_key in BANK_ORDER:
            item = (results or {}).get(bank_key)
            if item is None or BANK_NAMES[bank_key] in skipped_names:
                continue
            row_no = block_info["item_to_row"][BANK_NAMES[bank_key]]
            for field in WRITABLE_FIELDS:
                column = index_map.get(field)
                if column is None:
                    continue
                expected = as_number(item.get(field))
                if expected is None:
                    # summary 沒值的欄位可能已被 backfill 補上，不視為 mismatch
                    continue
                actual = as_number(ws.cell(row_no, column).value)
                equal = (
                    isinstance(expected, (int, float))
                    and isinstance(actual, (int, float))
                    and abs(float(expected) - float(actual)) <= 1e-9
                ) or expected == actual
                if not equal:
                    mismatches.append({
                        "item": BANK_NAMES[bank_key],
                        "field": FIELD_ALIASES[field][0],
                        "expected": expected,
                        "actual": actual,
                    })
            verified_banks.append(BANK_NAMES[bank_key])

        special_rows: dict[str, Any] = {}
        for item_name in SPECIAL_ITEMS:
            row_no = block_info["item_to_row"][item_name]
            non_blank = sum(
                1
                for field, column in index_map.items()
                if field not in ("yyyymm", "ad_year", "month_number", "rank", "bank", "item")
                and not is_blank(ws.cell(row_no, column).value)
            )
            entry: dict[str, Any] = {"row": row_no, "non_blank_fields": non_blank}
            if "avg_cards_per_person" in index_map:
                entry["avg_cards_per_person"] = ws.cell(row_no, index_map["avg_cards_per_person"]).value
            special_rows[item_name] = entry

        percent_anomalies = audit_block_percent_ratios(ws, index_map, block_info, sorted(PERCENT_DECIMAL_FIELDS))
        return {
            "status": "ok" if not mismatches and not percent_anomalies else "mismatch",
            "yyyymm": yyyymm,
            "block_start_row": block_info["start_row"],
            "block_rows": len(block_info["rows"]),
            "verified_banks": verified_banks,
            "mismatches": mismatches,
            "percent_anomalies": percent_anomalies,
            "special_rows": special_rows,
        }
    finally:
        wb.close()


def set_full_calc_on_load(wb) -> None:
    if hasattr(wb, "calculation"):
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            pass


def resolve_market_script_path(script_arg: str) -> Path:
    requested = Path(script_arg)
    if requested.exists():
        return requested
    sibling = Path(__file__).resolve().with_name(script_arg)
    if sibling.exists():
        return sibling
    raise FileNotFoundError(f"找不到市場總計腳本：{script_arg}")


def months_since_target(target_month: str) -> int:
    roc_year, month = parse_roc_month(target_month)
    ad_target_year = ad_year(roc_year)
    today = datetime.today()
    return max(0, (today.year - ad_target_year) * 12 + (today.month - month))


def effective_market_lookback_months(args, target_month: str | None) -> int:
    lookback = args.market_lookback_months
    if target_month:
        lookback = max(lookback, months_since_target(target_month) + 1)
    return lookback


def effective_jcic_backfill_months(args, target_month: str | None) -> int:
    lookback = args.jcic_backfill_months
    if target_month:
        lookback = max(lookback, months_since_target(target_month) + 1)
    return lookback


def run_market_total_script(args, workbook_path: Path, target_month: str | None = None, base_month: str | None = None) -> dict[str, Any]:
    script_path = resolve_market_script_path(args.market_script)
    cmd = [
        sys.executable,
        "-B",
        str(script_path),
        "--workbook",
        str(workbook_path),
        "--lookback-months",
        str(effective_market_lookback_months(args, target_month)),
    ]
    if base_month:
        cmd.extend(["--base-month", base_month])
    if args.market_overwrite:
        cmd.append("--overwrite")
    if args.market_insecure:
        cmd.append("--insecure")
    if args.market_zip_file:
        cmd.extend(["--zip-file", args.market_zip_file])
    for search_dir in args.market_zip_search_dir:
        cmd.extend(["--zip-search-dir", search_dir])
    if args.market_no_local_zip:
        cmd.append("--no-local-zip")
    if args.market_strict_workbook:
        cmd.append("--strict-workbook")

    delegate_started = time.perf_counter()
    completed = subprocess.run(cmd, capture_output=True)
    delegate_elapsed = round(time.perf_counter() - delegate_started, 1)
    stdout = decode_subprocess_output(completed.stdout).strip()
    stderr = decode_subprocess_output(completed.stderr).strip()

    if completed.returncode != 0:
        raise RuntimeError(
            "run_market_total.py 執行失敗："
            f"returncode={completed.returncode}; "
            f"stdout={stdout or '<empty>'}; stderr={stderr or '<empty>'}"
        )

    if stdout:
        try:
            payload: dict[str, Any] = json.loads(stdout)
        except json.JSONDecodeError:
            payload = {"raw_stdout": stdout}
    else:
        payload = {}

    payload.update({
        "delegated": True,
        "script": str(script_path),
        "command": cmd,
        "lookback_months": effective_market_lookback_months(args, target_month),
        "base_month": base_month,
        "elapsed_seconds": delegate_elapsed,
    })
    if stderr:
        payload["stderr"] = stderr
    return payload


def resolve_jcic_script_path(script_arg: str) -> Path:
    requested = Path(script_arg)
    if requested.exists():
        return requested
    sibling = Path(__file__).resolve().with_name(script_arg)
    if sibling.exists():
        return sibling
    raise FileNotFoundError(f"找不到 JCIC 腳本：{script_arg}")


def run_jcic_script(args, workbook_path: Path, base_month: str) -> dict[str, Any]:
    script_path = resolve_jcic_script_path(args.jcic_script)
    cmd = [
        sys.executable,
        "-B",
        str(script_path),
        "--workbook",
        str(workbook_path),
        "--sheet",
        MAIN_SHEET,
        "--base-month",
        base_month,
        "--backfill-months",
        str(effective_jcic_backfill_months(args, base_month)),
    ]
    if args.jcic_entry_url:
        cmd.extend(["--entry-url", args.jcic_entry_url])
    if args.jcic_source_url:
        cmd.extend(["--source-url", args.jcic_source_url])
    if args.jcic_label:
        cmd.extend(["--label", args.jcic_label])
    if args.jcic_overwrite:
        cmd.append("--overwrite")
    if args.jcic_insecure:
        cmd.append("--insecure")

    delegate_started = time.perf_counter()
    completed = subprocess.run(cmd, capture_output=True)
    delegate_elapsed = round(time.perf_counter() - delegate_started, 1)
    stdout = decode_subprocess_output(completed.stdout).strip()
    stderr = decode_subprocess_output(completed.stderr).strip()

    if completed.returncode != 0:
        raise RuntimeError(
            "jcic_avg_cards_update.py 執行失敗："
            f"returncode={completed.returncode}; "
            f"stdout={stdout or '<empty>'}; stderr={stderr or '<empty>'}"
        )

    if stdout:
        try:
            payload: dict[str, Any] = json.loads(stdout)
        except json.JSONDecodeError:
            payload = {"raw_stdout": stdout}
    else:
        payload = {}

    payload.update({
        "delegated": True,
        "script": str(script_path),
        "command": cmd,
        "base_month": base_month,
        "backfill_months": effective_jcic_backfill_months(args, base_month),
        "elapsed_seconds": delegate_elapsed,
    })
    if stderr:
        payload["stderr"] = stderr
    return payload


def resolve_bank_bureau_backfill_script_path(script_arg: str) -> Path:
    requested = Path(script_arg)
    if requested.exists():
        return requested
    sibling = Path(__file__).resolve().with_name(script_arg)
    if sibling.exists():
        return sibling
    raise FileNotFoundError(f"找不到銀行局逐家銀行回補腳本：{script_arg}")


def run_bank_bureau_backfill_script(args, workbook_path: Path, newest_month: str | None = None) -> dict[str, Any]:
    script_path = resolve_bank_bureau_backfill_script_path(args.bank_bureau_backfill_script)
    cmd = [
        sys.executable,
        "-B",
        str(script_path),
        "--workbook",
        str(workbook_path),
        "--lookback-months",
        str(args.bank_bureau_lookback_months),
    ]

    if args.bank_bureau_backfill_month:
        cmd.extend(["--month", args.bank_bureau_backfill_month])
    elif newest_month:
        # 自動模式：把 10 家銀行的資料月當作日曆缺口偵測的最新錨點，往前找缺漏年月
        cmd.extend(["--newest-month", newest_month])
    if args.bank_bureau_overwrite:
        cmd.append("--overwrite")
    if args.bank_bureau_insecure:
        cmd.append("--insecure")
    if args.bank_bureau_zip_file:
        cmd.extend(["--zip-file", args.bank_bureau_zip_file])
    for search_dir in args.bank_bureau_zip_search_dir:
        cmd.extend(["--zip-search-dir", search_dir])
    if args.bank_bureau_no_local_zip:
        cmd.append("--no-local-zip")
    if args.bank_bureau_strict_workbook:
        cmd.append("--strict-workbook")

    delegate_started = time.perf_counter()
    completed = subprocess.run(cmd, capture_output=True)
    delegate_elapsed = round(time.perf_counter() - delegate_started, 1)
    stdout = decode_subprocess_output(completed.stdout).strip()
    stderr = decode_subprocess_output(completed.stderr).strip()

    if completed.returncode != 0:
        raise RuntimeError(
            "run_bank_bureau_bank_backfill.py 執行失敗："
            f"returncode={completed.returncode}; "
            f"stdout={stdout or '<empty>'}; stderr={stderr or '<empty>'}"
        )

    if stdout:
        try:
            payload: dict[str, Any] = json.loads(stdout)
        except json.JSONDecodeError:
            payload = {"raw_stdout": stdout}
    else:
        payload = {}

    payload.update({
        "delegated": True,
        "script": str(script_path),
        "command": cmd,
        "elapsed_seconds": delegate_elapsed,
    })
    if stderr:
        payload["stderr"] = stderr
    return payload


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="更新銀行局信用卡公開資料.xlsx：以 summary JSON 寫入『歷史資料(年+月)』月 block，並委派市場總計、JCIC 與銀行局回補腳本"
    )
    parser.add_argument("--workbook", default=WORKBOOK_DEFAULT, help="Excel 檔案路徑，預設：銀行局信用卡公開資料.xlsx")
    parser.add_argument("--summary", default="", help="run_all_banks.py 產出的 JSON 檔；一般月更新模式必填")
    parser.add_argument("--target-month", "--month", dest="target_month", default="", help="目標月份，例如 115年06月；未指定時會嘗試由 summary 的 data_month 推定")
    parser.add_argument("--base-month", default="", help="JCIC / 市場總計回補的基期月份，例如 115年06月；未指定時會優先使用 target month")
    parser.add_argument("--allow-partial-summary", action="store_true", help="允許 summary 缺少部分銀行；預設為嚴格模式，缺任一銀行即報錯")

    parser.add_argument("--market-only", action="store_true", help="只更新銀行局市場總計，不寫入 10 家銀行月 block")
    parser.add_argument("--skip-market-total", action="store_true", help="不執行市場總計腳本")
    parser.add_argument("--market-script", default="run_market_total.py", help="市場總計腳本路徑，預設為與本檔同層的 run_market_total.py")
    parser.add_argument("--market-lookback-months", type=int, default=24, help="委派 run_market_total.py 時使用的 lookback months；若指定 target month 會自動提高到至少涵蓋該月份")
    parser.add_argument("--market-overwrite", action="store_true", help="委派 run_market_total.py 時加上 --overwrite")
    parser.add_argument("--market-insecure", action="store_true", help="委派 run_market_total.py 時加上 --insecure")
    parser.add_argument("--market-zip-file", default="", help="委派 run_market_total.py 時加上 --zip-file")
    parser.add_argument("--market-zip-search-dir", action="append", default=[], help="委派 run_market_total.py 時加上 --zip-search-dir，可重複指定")
    parser.add_argument("--market-no-local-zip", action="store_true", help="委派 run_market_total.py 時加上 --no-local-zip")
    parser.add_argument("--market-strict-workbook", action="store_true", help="委派 run_market_total.py 時加上 --strict-workbook")

    parser.add_argument("--skip-jcic-avg-cards", action="store_true", help="不執行 JCIC 腳本")
    parser.add_argument("--jcic-only", action="store_true", help="只更新 JCIC『平均每人持卡張數』；必須搭配 --base-month")
    parser.add_argument("--jcic-script", default="jcic_avg_cards_update.py", help="JCIC 腳本路徑，預設為與本檔同層的 jcic_avg_cards_update.py")
    parser.add_argument("--jcic-entry-url", default="", help="委派 jcic_avg_cards_update.py 時加上 --entry-url")
    parser.add_argument("--jcic-source-url", default="", help="委派 jcic_avg_cards_update.py 時加上 --source-url")
    parser.add_argument("--jcic-label", default="平均每人持卡張數", help="委派 jcic_avg_cards_update.py 時加上 --label")
    parser.add_argument("--jcic-backfill-months", type=int, default=24, help="委派 jcic_avg_cards_update.py 時往前回補空白時最多檢查幾個月份，預設 24")
    parser.add_argument("--jcic-overwrite", action="store_true", help="委派 jcic_avg_cards_update.py 時加上 --overwrite")
    parser.add_argument("--jcic-insecure", action="store_true", help="委派 jcic_avg_cards_update.py 時加上 --insecure")

    parser.add_argument("--skip-bank-bureau-backfill", action="store_true", help="不執行銀行局逐家銀行舊月份空白回補")
    parser.add_argument("--bank-bureau-backfill-script", default="run_bank_bureau_bank_backfill.py", help="銀行局逐家銀行回補腳本路徑，預設為與本檔同層的 run_bank_bureau_bank_backfill.py")
    parser.add_argument("--bank-bureau-backfill-month", default="", help="委派銀行局逐家銀行回補腳本時加上 --month；未指定時由該腳本自行決定回補範圍")
    parser.add_argument("--bank-bureau-lookback-months", type=int, default=24, help="委派銀行局逐家銀行回補腳本時使用的 lookback months，預設 24")
    parser.add_argument("--bank-bureau-overwrite", action="store_true", help="委派銀行局逐家銀行回補腳本時加上 --overwrite")
    parser.add_argument("--bank-bureau-insecure", action="store_true", help="委派銀行局逐家銀行回補腳本時加上 --insecure")
    parser.add_argument("--bank-bureau-zip-file", default="", help="委派銀行局逐家銀行回補腳本時加上 --zip-file")
    parser.add_argument("--bank-bureau-zip-search-dir", action="append", default=[], help="委派銀行局逐家銀行回補腳本時加上 --zip-search-dir，可重複指定")
    parser.add_argument("--bank-bureau-no-local-zip", action="store_true", help="委派銀行局逐家銀行回補腳本時加上 --no-local-zip")
    parser.add_argument("--bank-bureau-strict-workbook", action="store_true", help="委派銀行局逐家銀行回補腳本時加上 --strict-workbook")

    parser.add_argument("--annual-only", action="store_true", help="只執行年度整理；若未指定月份，會嘗試整理工作簿中所有完整年度")
    parser.add_argument("--skip-annual-sync", action="store_true", help="月資料更新完成後不自動整理年度資料")
    parser.add_argument("--audit-percent", action="store_true", help="只稽核、不寫入：掃描全表百分比欄（逾期比率、備抵呆帳提足率）的絕對範圍與市場總計 vs 十家銀行的相對關係，回報疑似 100 倍錯值")
    parser.add_argument("--repair-partial-blocks", action="store_true", help="偵測並刪除不完整的月 block 孤兒列（非 11 列一組）；只處理位於工作簿尾端、不影響其他列位置的 partial block，刪除前會在輸出 JSON 回報被移除的值")
    return parser


def main() -> None:
    configure_utf8_stdio()
    args = build_parser().parse_args()

    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到工作簿：{workbook_path}")

    output: dict[str, Any] = {
        "workbook": str(workbook_path),
        "sheet": MAIN_SHEET,
        "started_at": datetime.now().isoformat(timespec="seconds"),
    }
    warnings: list[str] = []
    timings: dict[str, float] = {}
    run_started = time.perf_counter()
    stage_started = run_started

    def mark(stage: str) -> None:
        nonlocal stage_started
        now = time.perf_counter()
        timings[stage] = round(now - stage_started, 1)
        stage_started = now

    explicit_target_month = roc_month_text(*parse_roc_month(args.target_month)) if args.target_month else None
    explicit_base_month = roc_month_text(*parse_roc_month(args.base_month)) if args.base_month else None

    if args.audit_percent:
        wb = load_workbook(workbook_path)
        ws = get_main_sheet(wb)
        index_map = header_index_map(ws, FIELD_ALIASES)
        require_headers(index_map, ["yyyymm", "rank", "item"], MAIN_SHEET)
        output["percent_audit"] = audit_percent_ratios(ws, index_map)
        if output["percent_audit"]["anomaly_count"]:
            warnings.append("百分比欄稽核發現異常格，請檢查 percent_audit.anomalies；修正方式：對該月用金管會 ZIP 跑 run_market_total.py --overwrite 或 backfill --overwrite")
        if warnings:
            output["warnings"] = warnings
        print(json.dumps(output, ensure_ascii=False, indent=2))
        return

    if args.repair_partial_blocks:
        wb = load_workbook(workbook_path)
        ws = get_main_sheet(wb)
        index_map = header_index_map(ws, FIELD_ALIASES)
        require_headers(index_map, ["yyyymm", "ad_year", "month_number", "rank", "item"], MAIN_SHEET)
        repair_report = repair_partial_blocks(ws, index_map)
        if repair_report["deleted_row_count"]:
            set_full_calc_on_load(wb)
            wb.save(workbook_path)
        output["repair_partial_blocks"] = repair_report
        if repair_report["needs_manual_review"]:
            warnings.append("部分不完整 block 無法自動修復（非表尾或非固定模板子集），請人工確認 needs_manual_review 清單")
        if warnings:
            output["warnings"] = warnings
        print(json.dumps(output, ensure_ascii=False, indent=2))
        return

    if args.annual_only:
        wb = load_workbook(workbook_path)
        ws = get_main_sheet(wb)
        index_map = header_index_map(ws, FIELD_ALIASES)
        require_headers(index_map, ["yyyymm", "ad_year", "month_number", "rank", "item", *WRITABLE_FIELDS, "avg_cards_per_person"], MAIN_SHEET)
        years: list[int]
        if explicit_target_month:
            roc_year_value, _ = parse_roc_month(explicit_target_month)
            years = [ad_year(roc_year_value)]
        else:
            years = sorted({normalize_year_value(ws.cell(r, index_map["yyyymm"]).value) for r in range(2, ws.max_row + 1) if normalize_year_value(ws.cell(r, index_map["yyyymm"]).value)}, reverse=True)
        annual_results = []
        for year_value in years:
            annual_results.append(auto_sync_annual_if_ready(ws, index_map, year_value))
        set_full_calc_on_load(wb)
        wb.save(workbook_path)
        output["annual_sync"] = annual_results if len(annual_results) != 1 else annual_results[0]
        if warnings:
            output["warnings"] = warnings
        print(json.dumps(output, ensure_ascii=False, indent=2))
        return

    if args.jcic_only:
        if not explicit_base_month:
            raise ValueError("--jcic-only 模式必須提供 --base-month")
        output["base_month"] = explicit_base_month
        output["jcic_avg_cards"] = run_jcic_script(args, workbook_path, explicit_base_month)
        if warnings:
            output["warnings"] = warnings
        print(json.dumps(output, ensure_ascii=False, indent=2))
        return

    results = None
    item_label_fix_summary: dict[str, Any] = {"changed_count": 0, "changed_rows": []}
    target_month = explicit_target_month
    summary_path = Path(args.summary) if args.summary else None

    if summary_path:
        if not summary_path.exists():
            raise FileNotFoundError(f"找不到 summary JSON：{summary_path}")
        results = load_summary_results(summary_path, allow_partial=args.allow_partial_summary)
        enrich_ctbc_debit_card_amount(results)
        validate_results_for_write(results, allow_partial=args.allow_partial_summary)
        if not target_month:
            target_month = infer_target_month_from_results(results)

    if not args.market_only and results is None:
        raise ValueError("一般月更新模式必須提供 --summary")

    backfill_base_month = explicit_base_month or determine_backfill_base_month(target_month, results)

    if not args.market_only:
        if not target_month:
            raise ValueError("無法判定 target month；請提供 --month，或讓 summary JSON 提供唯一 data_month")
        roc_year, month_number = parse_roc_month(target_month)

        wb = load_workbook(workbook_path)
        ws = get_main_sheet(wb)
        index_map = header_index_map(ws, FIELD_ALIASES)
        require_headers(index_map, ["yyyymm", "ad_year", "month_number", "rank", "item", *WRITABLE_FIELDS], MAIN_SHEET)

        partial_months = detect_partial_month_blocks(ws, index_map)
        if partial_months:
            raise ValueError(
                f"工作簿存在不完整的月 block（YYYYMM={', '.join(str(y) for y in partial_months)}）；"
                "請先執行 --repair-partial-blocks 修復後再跑月更新"
            )

        block_info = find_or_create_month_block(ws, index_map, ad_year(roc_year), month_number)
        refresh_block_metadata(ws, index_map, block_info, ad_year(roc_year), month_number)
        block_write_summary = write_bank_rows_to_block(ws, index_map, block_info, results or {}, target_month)
        skipped_banks = block_write_summary.get("skipped_banks") or []
        if skipped_banks:
            warnings.append(
                f"以下銀行資料月份與目標月 {target_month} 不符，已跳過（留待 backfill 補正確月份）：{', '.join(skipped_banks)}"
            )

        item_label_fix_summary = normalize_item_labels(ws, index_map)

        set_full_calc_on_load(wb)
        wb.save(workbook_path)
        mark("write_month_block")

        output.update({
            "target_month": target_month,
            "base_month": backfill_base_month,
            "yyyymm": roc_to_ad_yyyymm(roc_year, month_number),
            "block_created": block_info["created"],
            "block_start_row": block_info["start_row"],
            **block_write_summary,
        })
    else:
        output.update({
            "target_month": target_month,
            "base_month": backfill_base_month,
        })

    if not args.skip_jcic_avg_cards and not backfill_base_month and not args.market_only:
        raise ValueError("一般月更新模式若要更新 JCIC，請提供 --base-month，或提供可推定月份的 --summary / --month")

    # 委派順序：先把「更早月份的缺漏」用銀行局資料補齊（含 Top5/Top10 公式修復），
    # 讓所有資料列都完整後，再算跨列的衍生值（市場總計、JCIC 平均每人持卡張數），
    # 公式才不會建立在「還沒補完」的資料上。
    if not args.skip_bank_bureau_backfill and not args.market_only:
        output["bank_bureau_backfill"] = run_bank_bureau_backfill_script(args, workbook_path, target_month)
        mark("bank_bureau_backfill")

    if not args.skip_market_total:
        output["market_total"] = run_market_total_script(args, workbook_path, target_month, backfill_base_month)
        mark("market_total")

    if not args.skip_jcic_avg_cards:
        if not backfill_base_month:
            raise ValueError("無法判定 JCIC 回補基期月份")
        output["jcic_avg_cards"] = run_jcic_script(args, workbook_path, backfill_base_month)
        mark("jcic_avg_cards")

    if not args.skip_annual_sync and target_month:
        roc_year_value, _ = parse_roc_month(target_month)
        wb = load_workbook(workbook_path)
        ws = get_main_sheet(wb)
        index_map = header_index_map(ws, FIELD_ALIASES)
        require_headers(index_map, ["yyyymm", "ad_year", "month_number", "rank", "item", *WRITABLE_FIELDS, "avg_cards_per_person"], MAIN_SHEET)
        annual_sync_result = auto_sync_annual_if_ready(ws, index_map, ad_year(roc_year_value))
        set_full_calc_on_load(wb)
        wb.save(workbook_path)
        output["annual_sync"] = annual_sync_result
        mark("annual_sync")

    if not args.market_only and results is not None and target_month:
        output["verification"] = verify_written_block(workbook_path, target_month, results, skipped_banks)
        mark("verification")

    output["item_label_fix_summary"] = item_label_fix_summary

    # 月更新結尾全表稽核百分比欄：抓出任何來源/腳本改版造成的 100 倍錯值。
    if not args.market_only:
        wb = load_workbook(workbook_path)
        try:
            ws = get_main_sheet(wb)
            output["percent_audit"] = audit_percent_ratios(ws, header_index_map(ws, FIELD_ALIASES))
        finally:
            wb.close()
        if output["percent_audit"]["anomaly_count"]:
            warnings.append("百分比欄稽核發現異常格（percent_audit.anomalies），請回報使用者並用官方 ZIP --overwrite 修正該月")
        mark("percent_audit")

    timings["total"] = round(time.perf_counter() - run_started, 1)
    output["timings_seconds"] = timings
    output["finished_at"] = datetime.now().isoformat(timespec="seconds")

    if warnings:
        output["warnings"] = warnings

    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
