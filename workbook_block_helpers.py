from __future__ import annotations

"""`歷史資料(年+月)` 11 列 block 的共用讀寫工具（只在 writer 層 import）。"""

from copy import copy
from typing import Any, Callable

from openpyxl.utils import get_column_letter

from bank_aliases import (
    BANK_ITEM_ALIASES,
    BANK_NAMES,
    BLOCK_ITEMS,
    ITEM_RANKS,
    SPECIAL_ITEM_ALIASES,
    item_label,
    strip_item_prefix,
    year_block_key,
)

DEFAULT_FONT_NAME = '微軟正黑體'
BLOCK_LABEL = 'block'


def apply_default_font(cell, font_name: str = DEFAULT_FONT_NAME) -> None:
    new_font = copy(cell.font)
    new_font.name = font_name
    cell.font = new_font


def canonical_block_item(value: Any) -> str:
    """把 Item 欄各種寫法統一成簡稱：'01 中信'/'中國信託商業銀行' -> '中信'，'11 市場總計'/'市場總計(銀行局)' -> '市場總計'。"""
    text = strip_item_prefix(value)
    for special, aliases in SPECIAL_ITEM_ALIASES.items():
        if text == special or text in aliases:
            return special
    for bank_key, aliases in BANK_ITEM_ALIASES.items():
        if text in aliases:
            return BANK_NAMES[bank_key]
    return text


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith('=')


def copy_row_style(ws, src_row: int, dst_row: int) -> None:
    for column in range(1, ws.max_column + 1):
        src = ws.cell(src_row, column)
        dst = ws.cell(dst_row, column)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        apply_default_font(dst)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format
    if src_row in ws.row_dimensions:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
        ws.row_dimensions[dst_row].hidden = ws.row_dimensions[src_row].hidden


def copy_row_formulas(ws, src_row: int, dst_row: int) -> list[int]:
    """把樣板列的公式原樣複製到新列。公式用結構化參照 `[#This Row]` 與整欄參照，不需平移。"""
    copied: list[int] = []
    for column in range(1, ws.max_column + 1):
        value = ws.cell(src_row, column).value
        if is_formula(value):
            ws.cell(dst_row, column).value = value
            copied.append(column)
    return copied


def extend_table_ref(ws) -> dict[str, str]:
    """新增列後把工作表上的 Excel Table（如 `creditcard`）範圍延伸到最後一列，年度/月份/季度等結構化公式才會生效。"""
    changed: dict[str, str] = {}
    last = f"{get_column_letter(ws.max_column)}{ws.max_row}"
    for table in ws.tables.values():
        start = table.ref.split(':')[0]
        new_ref = f"{start}:{last}"
        if table.ref != new_ref:
            table.ref = new_ref
            if table.autoFilter is not None:
                table.autoFilter.ref = new_ref
            changed[table.displayName or table.name] = new_ref
    return changed


def set_row_value_if_present(ws, row_no: int, index_map: dict[str, int], field: str, value: Any) -> None:
    column = index_map.get(field)
    if column is not None:
        ws.cell(row_no, column).value = value


def set_block_row_metadata(
    ws,
    row_no: int,
    index_map: dict[str, int],
    *,
    yyyymm: Any,
    ad_year: int,
    month_number: int | str,
    item: str,
) -> None:
    """寫 YYYYMM / 排序編號 / Bank / Item。年度、月份欄若已是公式（Table 結構化公式）就不動。"""
    item = canonical_block_item(item)
    set_row_value_if_present(ws, row_no, index_map, 'yyyymm', yyyymm)
    for field, value in (('ad_year', ad_year), ('month_number', month_number)):
        column = index_map.get(field)
        if column is not None and not is_formula(ws.cell(row_no, column).value):
            ws.cell(row_no, column).value = value
    set_row_value_if_present(ws, row_no, index_map, 'rank', ITEM_RANKS.get(item))
    set_row_value_if_present(ws, row_no, index_map, 'bank', item)
    set_row_value_if_present(ws, row_no, index_map, 'item', item_label(item))


def validate_block_items(ws, index_map: dict[str, int], rows: list[int], key: Any, block_label: str = BLOCK_LABEL) -> dict[str, int]:
    if len(rows) != len(BLOCK_ITEMS):
        raise ValueError(f'YYYYMM={key} 的{block_label}列數不是 {len(BLOCK_ITEMS)}：實際 {len(rows)} 列')
    expected_rows = list(range(rows[0], rows[0] + len(BLOCK_ITEMS)))
    if rows != expected_rows:
        raise ValueError(f'YYYYMM={key} 的{block_label}不是連續 {len(BLOCK_ITEMS)} 列：{rows}')

    item_to_row: dict[str, int] = {}
    for offset, expected_item in enumerate(BLOCK_ITEMS):
        row_no = rows[offset]
        raw = ws.cell(row_no, index_map['item']).value
        if canonical_block_item(raw) != expected_item:
            raise ValueError(f'YYYYMM={key} 的{block_label}Item 不符合固定模板：row {row_no} 預期 {expected_item!r}，實際 {raw!r}')
        item_to_row[expected_item] = row_no
    return item_to_row


def find_block(ws, index_map: dict[str, int], key: Any, *, match: Callable[[Any], bool], block_label: str = BLOCK_LABEL) -> dict[str, Any] | None:
    rows = [row_no for row_no in range(2, ws.max_row + 1) if match(ws.cell(row_no, index_map['yyyymm']).value)]
    if not rows:
        return None
    item_to_row = validate_block_items(ws, index_map, rows, key, block_label)
    return {'yyyymm': key, 'start_row': rows[0], 'rows': rows, 'item_to_row': item_to_row, 'created': False}


def find_month_block(ws, index_map: dict[str, int], yyyymm: int, *, normalize_yyyymm: Callable[[Any], int | None], block_label: str = BLOCK_LABEL) -> dict[str, Any] | None:
    return find_block(ws, index_map, yyyymm, match=lambda v: normalize_yyyymm(v) == yyyymm, block_label=block_label)


def append_block(ws, index_map: dict[str, int], *, key: Any, ad_year: int, month_number: int | str) -> dict[str, Any]:
    """在工作表尾端新增一個 11 列 block：複製尾端 block 的樣式與公式，再寫入識別欄位，最後延伸 Table 範圍。"""
    old_max_row = ws.max_row
    start_row = old_max_row + 1
    template_start = old_max_row - len(BLOCK_ITEMS) + 1 if old_max_row >= len(BLOCK_ITEMS) + 1 else None
    item_to_row: dict[str, int] = {}
    rows: list[int] = []

    for offset, item in enumerate(BLOCK_ITEMS):
        row_no = start_row + offset
        if template_start:
            copy_row_style(ws, template_start + offset, row_no)
            copy_row_formulas(ws, template_start + offset, row_no)
        set_block_row_metadata(ws, row_no, index_map, yyyymm=key, ad_year=ad_year, month_number=month_number, item=item)
        item_to_row[item] = row_no
        rows.append(row_no)

    extend_table_ref(ws)
    return {'yyyymm': key, 'start_row': start_row, 'rows': rows, 'item_to_row': item_to_row, 'created': True}


def append_month_block(ws, index_map: dict[str, int], *, ad_year: int, month_number: int) -> dict[str, Any]:
    return append_block(ws, index_map, key=ad_year * 100 + month_number, ad_year=ad_year, month_number=month_number)


def append_year_block(ws, index_map: dict[str, int], *, ad_year: int) -> dict[str, Any]:
    return append_block(ws, index_map, key=year_block_key(ad_year), ad_year=ad_year, month_number='--')


def find_or_create_month_block(
    ws,
    index_map: dict[str, int],
    *,
    ad_year: int,
    month_number: int,
    normalize_yyyymm: Callable[[Any], int | None],
    block_label: str = BLOCK_LABEL,
) -> dict[str, Any]:
    found = find_month_block(ws, index_map, ad_year * 100 + month_number, normalize_yyyymm=normalize_yyyymm, block_label=block_label)
    return found or append_month_block(ws, index_map, ad_year=ad_year, month_number=month_number)


def refresh_block_metadata(ws, index_map: dict[str, int], block_info: dict[str, Any], *, ad_year: int, month_number: int | str) -> None:
    for item in BLOCK_ITEMS:
        set_block_row_metadata(ws, block_info['item_to_row'][item], index_map, yyyymm=block_info['yyyymm'], ad_year=ad_year, month_number=month_number, item=item)


def normalize_item_labels(ws, index_map: dict[str, int]) -> dict[str, Any]:
    """把所有 block 列的 Item 統一成「NN 名稱」、Bank 統一成簡稱；不屬於 block 模板的列不動。"""
    item_col = index_map.get('item')
    bank_col = index_map.get('bank')
    changed_rows: list[dict[str, Any]] = []
    if item_col is None:
        return {'changed_count': 0, 'changed_rows': changed_rows}
    for row_no in range(2, ws.max_row + 1):
        raw = ws.cell(row_no, item_col).value
        canonical = canonical_block_item(raw)
        if canonical not in ITEM_RANKS:
            continue
        label = item_label(canonical)
        if str(raw or '').strip() != label:
            ws.cell(row_no, item_col).value = label
            changed_rows.append({'row': row_no, 'from': raw, 'to': label})
        if bank_col is not None and str(ws.cell(row_no, bank_col).value or '').strip() != canonical:
            ws.cell(row_no, bank_col).value = canonical
    return {'changed_count': len(changed_rows), 'changed_rows': changed_rows}


def copy_formula_cells(ws, src_row: int, dst_row: int, columns: list[int], *, overwrite: bool) -> tuple[list[int], list[int]]:
    """把 src_row 指定欄的公式（含樣式）複製到 dst_row；回傳 (寫入欄, 因既有值略過欄)。"""
    written: list[int] = []
    skipped: list[int] = []
    for column in columns:
        source = ws.cell(src_row, column)
        if not is_formula(source.value):
            continue
        target = ws.cell(dst_row, column)
        if not overwrite and target.value not in (None, ''):
            skipped.append(column)
            continue
        target.value = source.value
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        apply_default_font(target)
        written.append(column)
    return written, skipped


def audit_block_percent_ratios(ws, index_map: dict[str, int], block_info: dict[str, Any], fields: list[str] | tuple[str, ...]) -> list[dict[str, Any]]:
    """
    寫入後稽核一個 block 的百分比欄：
    1. 每格套 percent_sanity_error 的絕對範圍。
    2. 市場總計列的逾期比率是全體發卡機構的加權平均，不可能比十家銀行最大值再高 10 倍以上，
       反過來也不可能低於十家銀行最小非零值的 1/10（這能抓到 100 倍錯值，即使數值本身仍在絕對範圍內）。
    回傳異常清單，空清單代表正常。
    """
    from percent_utils import percent_sanity_error  # 延後 import，避免與 percent_utils 互相依賴

    anomalies: list[dict[str, Any]] = []
    market_item = BLOCK_ITEMS[-1]
    bank_items = BLOCK_ITEMS[:-1]

    def numeric(row_no: int, column: int):
        value = ws.cell(row_no, column).value
        return float(value) if isinstance(value, (int, float)) else None

    for field in fields:
        column = index_map.get(field)
        if column is None:
            continue
        market_flagged = False
        for item, row_no in block_info['item_to_row'].items():
            reason = percent_sanity_error(field, ws.cell(row_no, column).value)
            if reason:
                anomalies.append({'yyyymm': block_info['yyyymm'], 'row': row_no, 'item': item, 'field': field, 'value': ws.cell(row_no, column).value, 'reason': reason})
                market_flagged = market_flagged or item == market_item
        if not field.startswith('overdue_') or market_flagged:
            continue
        market = numeric(block_info['item_to_row'][market_item], column)
        banks = [v for v in (numeric(block_info['item_to_row'][item], column) for item in bank_items) if v is not None]
        if market is None or not banks:
            continue
        bank_max = max(banks)
        bank_min_nonzero = min([v for v in banks if v > 0], default=None)
        if bank_max > 0 and market > 10 * bank_max:
            anomalies.append({'yyyymm': block_info['yyyymm'], 'row': block_info['item_to_row'][market_item], 'item': market_item, 'field': field, 'value': market, 'reason': f'市場總計 {market} 超過十家銀行最大值 {bank_max} 的 10 倍，疑似 100 倍錯值'})
        elif bank_max == 0 and market > 0.001:
            anomalies.append({'yyyymm': block_info['yyyymm'], 'row': block_info['item_to_row'][market_item], 'item': market_item, 'field': field, 'value': market, 'reason': f'十家銀行皆為 0 但市場總計 {market} 超過 0.1%，疑似 100 倍錯值'})
        elif bank_min_nonzero is not None and 0 < market < bank_min_nonzero / 10:
            anomalies.append({'yyyymm': block_info['yyyymm'], 'row': block_info['item_to_row'][market_item], 'item': market_item, 'field': field, 'value': market, 'reason': f'市場總計 {market} 低於十家銀行最小非零值 {bank_min_nonzero} 的 1/10，疑似被多除 100'})
    return anomalies
