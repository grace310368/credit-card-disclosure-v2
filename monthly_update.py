#!/usr/bin/env python3
from __future__ import annotations

"""GitHub Actions 用的月更新驅動程式：串起 SKILL.md 情境 A 的所有步驟，不需要 Claude。

輸出（都在 --out-dir，預設 out/）：
  groupA.json / groupB.json / firstbank.json / retry_<bank>.json / summary.json / update.json
  status.json  — 給 workflow 與通知用：target_month、yyyymm、branch、成功銀行數、verification、audit、warnings
  pr_body.md   — PR 內容
並產生 reports/monthly-update-<yyyymm>.md。任一關鍵步驟失敗以非 0 結束，status.json 仍會寫出 error。
"""

import argparse
import json
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

sys.dont_write_bytecode = True
BASE_DIR = Path(__file__).resolve().parent
GROUPS = [
    ("groupA", "ctbc,fubon,cathay,esun,taishin,feib", 95, []),
    ("groupB", "dbs,ubot,sinopac", 120, []),
    ("firstbank", "firstbank", 250, ["--allow-insecure"]),
]
BANK_NAMES = {"ctbc": "中信", "fubon": "富邦", "cathay": "國泰", "esun": "玉山", "taishin": "台新",
              "dbs": "星展", "ubot": "聯邦", "sinopac": "永豐", "firstbank": "第一", "feib": "遠東"}


def run(cmd: list[str], timeout: int) -> subprocess.CompletedProcess:
    print("$", " ".join(cmd), flush=True)
    return subprocess.run(cmd, cwd=BASE_DIR, capture_output=True, text=True, timeout=timeout)


def latest_yyyymm(workbook: Path) -> int:
    from openpyxl import load_workbook
    ws = load_workbook(workbook, read_only=True)["歷史資料(年+月)"]
    values = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    return max(v for v in values if isinstance(v, int))


def yyyymm_to_roc(yyyymm: int) -> str:
    y, m = divmod(yyyymm, 100)
    return f"{y - 1911}年{m:02d}月"


def roc_to_yyyymm(text: str) -> int | None:
    m = re.fullmatch(r"(\d{3})年(\d{2})月", text or "")
    return (int(m.group(1)) + 1911) * 100 + int(m.group(2)) if m else None


def next_month(yyyymm: int) -> int:
    y, m = divmod(yyyymm, 100)
    return y * 100 + m + 1 if m < 12 else (y + 1) * 100 + 1


def load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default="銀行局信用卡公開資料.xlsx")
    ap.add_argument("--month", default="", help="指定目標月（民國，如 115年09月）；預設為工作簿最新月的下一個月")
    ap.add_argument("--out-dir", default="out")
    args = ap.parse_args()

    out = BASE_DIR / args.out_dir
    out.mkdir(exist_ok=True)
    workbook = BASE_DIR / args.workbook
    status: dict[str, Any] = {"started_at": datetime.now().isoformat(timespec="seconds"), "workbook": args.workbook}
    t0 = time.perf_counter()

    def finish(code: int, **extra: Any) -> int:
        status.update(extra)
        status["elapsed_seconds"] = round(time.perf_counter() - t0, 1)
        status["finished_at"] = datetime.now().isoformat(timespec="seconds")
        (out / "status.json").write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({k: v for k, v in status.items() if k != "pr_body"}, ensure_ascii=False, indent=2))
        return code

    if not workbook.exists():
        return finish(1, error=f"找不到工作簿 {workbook}")

    latest = latest_yyyymm(workbook)
    target = roc_to_yyyymm(args.month) if args.month else next_month(latest)
    if target is None:
        return finish(1, error=f"--month 格式錯誤：{args.month}")
    status.update(latest_yyyymm=latest, requested_yyyymm=target)
    target_roc = yyyymm_to_roc(target)

    def find_ctbc_file(yyyymm: int) -> Path | None:
        return next((p for p in [BASE_DIR / f"中信信用卡{yyyymm}.xlsx", BASE_DIR / "input" / f"中信信用卡{yyyymm}.xlsx"] if p.exists()), None)

    ctbc_file = find_ctbc_file(target)

    # ---- 抓取 ----
    json_files: list[str] = []
    for name, banks, timeout, extra in GROUPS:
        cmd = [sys.executable, "-B", "run_all_banks.py", "--banks", banks, "--month", target_roc, "--timeout", str(timeout), *extra,
               "--json-output", f"{args.out_dir}/{name}.json"]
        proc = run(cmd, timeout=timeout * len(banks.split(",")) + 60)
        if proc.returncode != 0 and not (out / f"{name}.json").exists():
            return finish(1, error=f"run_all_banks {name} 失敗：{proc.stderr[-500:]}")
        json_files.append(f"{args.out_dir}/{name}.json")

    proc = run([sys.executable, "-B", "run_all_banks.py", "--collect-only", *json_files, "--json-output", f"{args.out_dir}/summary.json"], 120)
    summary = load(out / "summary.json")
    failed = [r["bank_key"] for r in summary["results"] if r.get("status") == "failed"]
    retries: list[str] = []
    for bank in failed:  # 情境 B：失敗的銀行補跑一次
        extra = ["--allow-insecure"] if bank == "firstbank" else []
        run([sys.executable, "-B", "run_all_banks.py", "--banks", bank, "--month", target_roc, "--timeout", "250", *extra,
             "--json-output", f"{args.out_dir}/retry_{bank}.json"], 320)
        if (out / f"retry_{bank}.json").exists():
            retries.append(f"{args.out_dir}/retry_{bank}.json")
    if retries:
        run([sys.executable, "-B", "run_all_banks.py", "--collect-only", f"{args.out_dir}/summary.json", *retries, "--json-output", f"{args.out_dir}/summary.json"], 120)
        summary = load(out / "summary.json")

    results = summary["results"]
    bank_status = {r["bank_key"]: (r.get("status"), r.get("data_month")) for r in results}
    ok = [k for k, (s, _) in bank_status.items() if s == "success"]
    status.update(bank_status={BANK_NAMES.get(k, k): f"{s}（{m}）" for k, (s, m) in bank_status.items()},
                  success_count=len(ok), failed_banks=[BANK_NAMES.get(k, k) for k in bank_status if bank_status[k][0] == "failed"],
                  retried_banks=[BANK_NAMES.get(b, b) for b in failed])
    if not ok:
        return finish(1, error="10 家銀行全部抓取失敗")

    # 官網尚未發布目標月時，改以 summary 的最新月為目標
    months = [roc_to_yyyymm(str(r.get("data_month") or "")) for r in results if r.get("status") == "success"]
    months = [m for m in months if m]
    actual = max(months) if months else target
    if actual != target:
        status["note"] = f"官網尚未發布 {target_roc}，各行最新資料月為 {yyyymm_to_roc(actual)}，改以該月為目標"
        target, target_roc = actual, yyyymm_to_roc(actual)
    ctbc_file = find_ctbc_file(target)  # 目標月可能已改為官網實際月份，重新確認中信檔
    status.update(target_month=target_roc, yyyymm=target, branch=f"auto/monthly-update-{target}", ctbc_file=ctbc_file.name if ctbc_file else None)
    if actual <= latest:
        status["note"] = (status.get("note", "") + f"；工作簿已有 {target}，本次為同月重寫").strip("；")

    # ---- 回寫 ----
    partial = len(ok) < 10 or ctbc_file is None
    cmd = [sys.executable, "-B", "update_credit_card_workbook.py", "--workbook", args.workbook, "--summary", f"{args.out_dir}/summary.json"]
    if partial:
        cmd.append("--allow-partial-summary")
    proc = None
    for attempt in range(1, 3):  # 金管會 / JCIC 偶發連線錯誤，重跑一次即可
        proc = run(cmd, timeout=900)
        if proc.returncode == 0:
            break
        status[f"update_attempt_{attempt}_error"] = proc.stderr[-400:]
        time.sleep(20)
    (out / "update.json").write_text(proc.stdout, encoding="utf-8")
    (out / "update.err").write_text(proc.stderr, encoding="utf-8")
    if proc.returncode != 0:
        return finish(1, error=f"update_credit_card_workbook 失敗（已重試）：{proc.stderr[-600:]}")
    update = json.loads(proc.stdout)
    verification = update.get("verification") or {}
    audit = update.get("percent_audit") or {}
    jcic = update.get("jcic_avg_cards") or {}
    market = update.get("market_total") or {}
    status.update(verification=verification.get("status"), mismatches=len(verification.get("mismatches") or []),
                  percent_anomalies=audit.get("anomaly_count"), warnings=update.get("warnings") or [],
                  jcic_latest=jcic.get("latest_available_month"),
                  market_months=[r.get("ad_yyyymm") for r in market.get("results", []) if r.get("action", "").endswith("update") or r.get("action", "").startswith("backfill")],
                  timings=update.get("timings_seconds"))

    # ---- 報告與 PR 內容 ----
    report = BASE_DIR / "reports" / f"monthly-update-{target}.md"
    run([sys.executable, "-B", "run_report.py", "--summary", f"{args.out_dir}/summary.json", "--update-output", f"{args.out_dir}/update.json",
         *sum([["--retry", r] for r in retries], []), "--label", f"{target_roc}（GitHub Actions）", "--output", str(report)], 120)
    report_text = report.read_text(encoding="utf-8") if report.exists() else ""
    review = report_text.split("## 檢討與建議", 1)[-1].strip() if "## 檢討與建議" in report_text else "（報告產生失敗）"

    lines = ["## 結果摘要", "",
             f"- 目標月：{target_roc}（{target}）" + (f"，{status['note']}" if status.get("note") else ""),
             f"- 銀行抓取：成功 {len(ok)}／10" + (f"，失敗 {status['failed_banks']}" if status["failed_banks"] else "") + (f"，補跑 {status['retried_banks']}" if failed else ""),
             f"- 中信檔：{'已就位（' + ctbc_file.name + '）' if ctbc_file else '缺少，當月轉帳卡簽帳金額留空'}",
             f"- verification：{verification.get('status')}，mismatch {status['mismatches']} 格；percent_audit：{audit.get('anomaly_count')} 格異常（{audit.get('blocks_audited')} 個 block）",
             f"- 市場總計寫入月份：{status['market_months'] or '無（金管會尚未發布）'}；JCIC 最新可用：{status['jcic_latest']}",
             f"- 警告：{status['warnings'] or '無'}",
             "", "## 檢討與建議", "", review, "",
             f"完整報告：`reports/monthly-update-{target}.md`", "",
             "## 待人工處理", "", "- 檢查後合併本 PR", "",
             "🤖 由 GitHub Actions 自動排程產生"]
    (out / "pr_body.md").write_text("\n".join(lines), encoding="utf-8")
    return finish(0, report=str(report.relative_to(BASE_DIR)))


if __name__ == "__main__":
    raise SystemExit(main())
